"""Regresja formatów wyjścia: eksport we wszystkich formatach i zapis projektu.

Projekt z poprzedniego wydania (``data/``) jest otwierany w oknie programu
i eksportowany przez ``ExportController`` do każdego formatu z rejestru,
a potem zapisywany jako projekt. Wynik porównujemy z plikami w ``golden/``.
Formaty tekstowe są porównywane w całości (bez daty zapisu i ścieżek tego
komputera); DOCX, ODT i PDF - jako czytelny zapis treści: akapity, style,
indeks górny (``^{…}``), tabela metryczki i właściwości pliku.
"""

import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest
from docx import Document as DocxDocument
from PySide6.QtPdf import QPdfDocument
from PySide6.QtWidgets import QFileDialog
from tests.e2e.conftest import DATA_DIR, REPO_ROOT, Golden

from transkryptor.export.options import ExportOptions
from transkryptor.export.registry import EXPORTERS, exporter_by_key
from transkryptor.ui.main_window import MainWindow

PROJECT = DATA_DIR / "wywiad-v0.6.0.transkr"

ODF_TEXT = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
ODF_STYLE = "urn:oasis:names:tc:opendocument:xmlns:style:1.0"
ODF_TABLE = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"


@pytest.fixture
def window(qtbot, add_window) -> MainWindow:
    win = add_window(MainWindow())
    assert win.session.open_project(PROJECT)
    qtbot.waitUntil(lambda: win.player_bar.has_media, timeout=10000)
    return win


def save_to(monkeypatch, path: Path) -> None:
    monkeypatch.setattr(
        QFileDialog,
        "getSaveFileName",
        staticmethod(lambda *args, **kwargs: (str(path), "")),
    )


def machine_independent(text: str, tmp_path: Path) -> str:
    """Bez ścieżek katalogu tymczasowego i repozytorium tego komputera."""
    for path, name in ((tmp_path, "<tmp>"), (REPO_ROOT, "<repo>")):
        text = text.replace(path.as_uri(), f"file://{name}")
        text = text.replace(str(path), name)
    return text


def mark_superscript(text: str, superscript: bool) -> str:
    return f"^{{{text}}}" if superscript and text else text


def docx_summary(path: Path) -> str:
    docx = DocxDocument(str(path))
    props = docx.core_properties
    lines = [
        f"author: {props.author}",
        f"title: {props.title}",
        f"subject: {props.subject}",
        f"keywords: {props.keywords}",
        f"comments: {props.comments}",
        "",
    ]
    for table in docx.tables:
        for row in table.rows:
            lines.append("| " + " | ".join(cell.text for cell in row.cells) + " |")
        lines.append("")
    for paragraph in docx.paragraphs:
        text = "".join(
            mark_superscript(run.text, bool(run.font.superscript))
            for run in paragraph.runs
        )
        lines.append(f"[{paragraph.style.name if paragraph.style else ''}] {text}")
    return "\n".join(lines) + "\n"


def odt_summary(path: Path) -> str:
    """Treść content.xml; nazwy stylów Qt zmieniają się między przebiegami."""
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("content.xml"))
    superscript = {
        style.get(f"{{{ODF_STYLE}}}name")
        for style in root.iter(f"{{{ODF_STYLE}}}style")
        if any(
            props.get(f"{{{ODF_STYLE}}}text-position", "").startswith("super")
            for props in style.iter(f"{{{ODF_STYLE}}}text-properties")
        )
    }

    def paragraph_text(element: ElementTree.Element) -> str:
        parts = [element.text or ""]
        for child in element:
            styled = child.get(f"{{{ODF_TEXT}}}style-name") in superscript
            parts.append(mark_superscript(paragraph_text(child), styled))
            parts.append(child.tail or "")
        return "".join(parts)

    lines: list[str] = []
    for element in root.iter():
        if element.tag == f"{{{ODF_TABLE}}}table-row":
            cells = [
                "".join(paragraph_text(p) for p in cell.iter(f"{{{ODF_TEXT}}}p"))
                for cell in element.iter(f"{{{ODF_TABLE}}}table-cell")
            ]
            lines.append("| " + " | ".join(cells) + " |")
        elif element.tag == f"{{{ODF_TEXT}}}p" and not _in_table(root, element):
            lines.append(paragraph_text(element))
    return "\n".join(lines) + "\n"


def _in_table(root: ElementTree.Element, paragraph: ElementTree.Element) -> bool:
    return any(
        paragraph in cell.iter(f"{{{ODF_TEXT}}}p")
        for cell in root.iter(f"{{{ODF_TABLE}}}table-cell")
    )


def pdf_summary(path: Path) -> str:
    """Liczba stron i tekst stron (łamanie wierszy zależy od fontów systemu)."""
    pdf = QPdfDocument()
    assert pdf.load(str(path)) == QPdfDocument.Error.None_
    try:
        pages = [
            " ".join(pdf.getAllText(page).text().split())
            for page in range(pdf.pageCount())
        ]
    finally:
        pdf.close()
    return f"pages: {len(pages)}\n" + "".join(
        f"\n--- {number} ---\n{text}\n" for number, text in enumerate(pages, 1)
    )


def text_summary(path: Path, tmp_path: Path) -> str:
    text = path.read_text(encoding="utf-8")
    # ELAN: data zapisu i ścieżka nagrania względem katalogu eksportu.
    text = re.sub(r' DATE="[^"]*"', ' DATE="<date>"', text)
    text = re.sub(r' RELATIVE_MEDIA_URL="[^"]*"', ' RELATIVE_MEDIA_URL="<rel>"', text)
    return machine_independent(text, tmp_path)


SUMMARIES = {"docx": docx_summary, "odt": odt_summary, "pdf": pdf_summary}


@pytest.mark.parametrize("key", [exporter.key for exporter in EXPORTERS])
def test_export_matches_golden(
    key, window, monkeypatch, tmp_path, golden: Golden, no_blocking_dialogs
) -> None:
    exporter = exporter_by_key(key)
    path = tmp_path / f"wywiad{exporter.suffix}"
    save_to(monkeypatch, path)

    assert window.exporter.export_with(exporter, ExportOptions(), False)

    summary = SUMMARIES.get(key)
    content = summary(path) if summary else text_summary(path, tmp_path)
    report = window.exporter.last_report
    warnings = report.warnings if report is not None else []
    golden(
        f"export/wywiad.{key}.txt",
        content + "".join(f"\n=== raport ===\n{line}\n" for line in warnings),
    )
    errors = [call for call in no_blocking_dialogs if call[0] == "critical"]
    assert not errors


def test_anonymized_docx_matches_golden(
    window, monkeypatch, tmp_path, golden: Golden
) -> None:
    path = tmp_path / "wywiad-anonim.docx"
    save_to(monkeypatch, path)
    window.export_anonymized_action.trigger()
    golden("export/wywiad-anonim.docx.txt", docx_summary(path))
    assert window.document.is_dirty is False  # projekt był wyeksportowany w 0.6.0


def test_saved_project_matches_golden(
    window, monkeypatch, tmp_path, golden: Golden
) -> None:
    """Bieżący format pliku projektu po migracji z 0.6.0."""
    path = tmp_path / "wywiad.transkr"
    save_to(monkeypatch, path)
    assert window.session.save_project_as()

    data = json.loads(path.read_text(encoding="utf-8"))
    audio = data["audio"]
    assert Path(audio["absolute_path"]).samefile(
        REPO_ROOT / "tests" / "fixtures" / "audio" / "sample.mp3"
    )
    audio["absolute_path"] = "<repo>/tests/fixtures/audio/sample.mp3"
    audio["relative_path"] = "<rel>"
    golden(
        "project/wywiad.transkr.json",
        json.dumps(data, ensure_ascii=False, indent=1, sort_keys=True) + "\n",
    )
