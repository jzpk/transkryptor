"""Komentarze recenzji jako komentarze Worda w DOCX (faza 14, REQ-27)."""

from docx import Document as DocxReader

from transkryptor.document.comments import Comment, Reply
from transkryptor.document.model import Document
from transkryptor.export.docx_export import export_docx
from transkryptor.export.docx_import import import_docx
from transkryptor.export.options import ExportOptions

TEXT = "tak było\npamientam dobrze"
COMMENT = Comment(
    id="c1",
    start=4,
    end=8,
    author="Ewa Kowalska",
    created="2026-10-24T14:05",
    text="Raczej „buło”?",
    quote="było",
    replies=(Reply("Jan", "2026-10-24T15:00", "Poprawione."),),
)


def commented_text(read, comment_id: int) -> str:
    """Tekst objęty komentarzem (między znacznikami zakresu)."""
    body = read.element.body
    inside = False
    pieces: list[str] = []
    for element in body.iter():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "commentRangeStart" and element.get(_w("id")) == str(comment_id):
            inside = True
        elif tag == "commentRangeEnd" and element.get(_w("id")) == str(comment_id):
            break
        elif inside and tag == "t":
            pieces.append(element.text or "")
    return "".join(pieces)


def _w(name: str) -> str:
    return f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{name}"


def document() -> Document:
    doc = Document(text=TEXT)
    doc.superscript_ranges = [(11, 12)]
    return doc


def test_comment_with_reply_covers_its_fragment(tmp_path) -> None:
    target = tmp_path / "recenzja.docx"
    export_docx(document(), target, comments=[COMMENT])
    read = DocxReader(str(target))
    (comment,) = list(read.comments)
    assert comment.author == "Ewa Kowalska"
    assert comment.initials == "EK"
    assert [p.text for p in comment.paragraphs] == [
        "Raczej „buło”?",
        "Odpowiedź — Jan (2026-10-24 15:00): Poprawione.",
    ]
    assert commented_text(read, comment.comment_id) == "było"
    assert comment._comment_elm.get(_w("date")) == "2026-10-24T14:05:00"
    # Podział runów na granicach komentarza nie zmienia tekstu ani formatu.
    paragraph = read.paragraphs[-1]
    assert paragraph.text == "pamientam dobrze"
    assert [r.text for r in paragraph.runs if r.font.superscript] == ["m"]


def test_comment_across_paragraphs_and_point_comment(tmp_path) -> None:
    across = Comment("c2", 4, 18, "A", "", "dwa akapity")
    point = Comment("c3", 9, 9, "", "", "usunięty fragment", quote="x")
    target = tmp_path / "recenzja.docx"
    export_docx(document(), target, comments=[across, point])
    read = DocxReader(str(target))
    first, second = list(read.comments)
    assert commented_text(read, first.comment_id) == "byłopamientam"
    assert second.author == "Bez podpisu"
    assert commented_text(read, second.comment_id)


def test_anonymized_export_omits_comments(tmp_path) -> None:
    target = tmp_path / "anonim.docx"
    export_docx(document(), target, anonymize=True, comments=[COMMENT])
    assert list(DocxReader(str(target)).comments) == []


def test_comments_on_template(tmp_path) -> None:
    template = tmp_path / "szablon.docx"
    docx = DocxReader()
    docx.add_paragraph("Instytut")
    docx.add_paragraph("{{transkrypcja}}", style="Quote")
    docx.save(str(template))
    target = tmp_path / "wynik.docx"
    export_docx(
        document(),
        target,
        options=ExportOptions(template_path=str(template)),
        comments=[COMMENT],
    )
    read = DocxReader(str(target))
    (comment,) = list(read.comments)
    assert commented_text(read, comment.comment_id) == "było"


def test_import_ignores_comments(tmp_path) -> None:
    target = tmp_path / "recenzja.docx"
    export_docx(document(), target, comments=[COMMENT])
    imported = import_docx(target)
    assert imported.text == TEXT
    assert imported.superscript_ranges == [(11, 12)]
