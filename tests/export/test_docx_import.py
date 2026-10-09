"""Testy importu DOCX (faza 08, propozycja 17, ACC-33)."""

import pytest
from docx import Document as DocxDocument

from transkryptor.document.metadata import DEFAULT_FIELDS, new_custom_field
from transkryptor.document.model import Document
from transkryptor.errors import ImportDocxError
from transkryptor.export.docx_export import export_docx
from transkryptor.export.docx_import import import_docx


def test_round_trip_of_exported_document(tmp_path) -> None:
    """ACC-33: tekst, indeks górny i metryczka równe stanowi sprzed eksportu."""
    custom = new_custom_field("Numer taśmy")
    fields = (*DEFAULT_FIELDS, custom)
    doc = Document(
        text="uod tego czasu … som\npamientam\tjeż…",
        author="Łukasz Żółw",
        date="2026-10-09",
        metadata={
            "signature": "AdK_1954",
            "place": "Ocieszyn",
            "informant": "KA",
            custom.key: "T-12",
        },
    )
    doc.superscript_ranges = [(0, 1), (19, 20), (25, 26)]
    path = tmp_path / "eksport.docx"
    export_docx(doc, path, fields)

    imported = import_docx(path, fields)

    assert imported.text == doc.text
    assert imported.superscript_ranges == doc.superscript_ranges
    assert imported.author == doc.author
    assert imported.date == doc.date
    assert imported.metadata == doc.metadata
    assert imported.report.is_lossless


def test_round_trip_of_empty_document(tmp_path) -> None:
    path = tmp_path / "pusty.docx"
    export_docx(Document(), path)
    imported = import_docx(path)
    assert imported.text == ""
    assert imported.superscript_ranges == []


def test_report_lists_skipped_elements(tmp_path) -> None:
    docx = DocxDocument()
    first = docx.add_paragraph()
    first.add_run("pierwszy ")
    first.add_run("gruby").bold = True
    second = docx.add_paragraph()
    second.add_run("bendzie")
    second.runs[0].font.superscript = False
    up = second.add_run("n")
    up.font.superscript = True
    docx.add_table(rows=1, cols=3)
    docx.add_paragraph("trzeci")
    path = tmp_path / "obcy.docx"
    docx.save(str(path))

    imported = import_docx(path)

    assert imported.text == "pierwszy gruby\nbendzien\ntrzeci"
    assert imported.superscript_ranges == [(22, 23)]
    report = imported.report
    assert report.merged_paragraphs == 3
    assert report.formatted_runs == {"pogrubienie": 1}
    assert report.skipped_tables == 1
    lines = report.lines()
    assert any("3 akapity" in line for line in lines)
    assert any("pogrubienie (1)" in line for line in lines)


def test_unknown_metadata_rows_are_reported(tmp_path) -> None:
    docx = DocxDocument()
    table = docx.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Sygnatura:"
    table.rows[0].cells[1].text = "X_1"
    table.rows[1].cells[0].text = "Pole zespołu"
    table.rows[1].cells[1].text = "wartość"
    docx.add_paragraph("tekst")
    path = tmp_path / "metryczka.docx"
    docx.save(str(path))

    imported = import_docx(path)

    assert imported.metadata == {"signature": "X_1"}
    assert imported.report.unknown_metadata == ("Pole zespołu",)
    assert imported.report.skipped_tables == 0


def test_not_a_docx_raises_import_error(tmp_path) -> None:
    path = tmp_path / "zly.docx"
    path.write_text("to nie jest docx")
    with pytest.raises(ImportDocxError) as info:
        import_docx(path)
    assert "zly.docx" in info.value.user_message
    assert info.value.retry_hint


def test_reads_external_transcription_layout(tmp_path) -> None:
    """ACC-33: układ ręcznej transkrypcji spoza aplikacji (jak AdK_1954.docx).

    Nagłówek pogrubiony i podkreślony, akapity kwestionariusza, kursywa
    i indeks górny ustawiony stylem znakowym. Sam plik ``AdK_1954.docx`` jest
    lokalnym nagraniem badawczym (poza repozytorium) — sprawdzany ręcznie.
    """
    from docx.enum.style import WD_STYLE_TYPE

    docx = DocxDocument()
    style = docx.styles.add_style("Nosowość", WD_STYLE_TYPE.CHARACTER)
    style.font.superscript = True
    heading = docx.add_paragraph().add_run("CHRZEST")
    heading.bold = True
    heading.underline = True
    docx.add_paragraph()
    line = docx.add_paragraph("- przedsie")
    line.add_run("m", style="Nosowość")
    line.add_run("biorstwo ")
    line.add_run("(?)").italic = True
    path = tmp_path / "kwestionariusz.docx"
    docx.save(str(path))

    imported = import_docx(path)

    assert imported.text == "CHRZEST\n\n- przedsiembiorstwo (?)"
    assert [imported.text[s:e] for s, e in imported.superscript_ranges] == ["m"]
    assert imported.report.formatted_runs == {
        "pogrubienie": 1,
        "podkreślenie": 1,
        "kursywa": 1,
    }
