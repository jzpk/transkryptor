"""Testy szablonu eksportu DOCX (faza 12, ACC-55, ACC-56)."""

import struct
import zlib

from docx import Document as DocxDocument
from docx.shared import Cm

from transkryptor.document.metadata import DEFAULT_FIELDS, new_custom_field
from transkryptor.document.model import Document
from transkryptor.export.docx_export import export_docx
from transkryptor.export.docx_import import import_docx
from transkryptor.export.docx_template import (
    check_template,
    placeholder_key,
    placeholder_names,
)
from transkryptor.export.options import ExportOptions


def png_bytes() -> bytes:
    """Najmniejszy poprawny PNG 1×1 (logo w nagłówku szablonu)."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    pixels = zlib.compress(b"\x00\xff\x00\x00")
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", pixels)
        + chunk(b"IEND", b"")
    )


def make_template(tmp_path, transcript=True):
    logo = tmp_path / "logo.png"
    logo.write_bytes(png_bytes())
    docx = DocxDocument()
    section = docx.sections[0]
    section.left_margin = Cm(3.5)
    header = section.header.paragraphs[0]
    header.add_run().add_picture(str(logo), width=Cm(1))
    header.add_run(" Sygnatura: {{syg")  # placeholder rozbity na runy
    header.add_run("natura}}")
    section.footer.paragraphs[0].text = "Badacz: {{ Badacz }} · {{nieznane}}"
    docx.add_paragraph("Instytut Języka Polskiego")
    table = docx.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Informator"
    table.rows[0].cells[1].text = "{{informator}} ({{Miejscowość}})"
    if transcript:
        docx.add_paragraph("{{transkrypcja}}", style="Quote")
    docx.add_paragraph("Koniec dokumentu")
    path = tmp_path / "szablon.docx"
    docx.save(str(path))
    return path


def make_document() -> Document:
    document = Document(
        text="Rok urodzenia?\nTysionc dziewińćset",
        author="Ewa Nowak",
        date="2026-10-18",
        metadata={
            "signature": "AdK_1954",
            "researcher": "J. Kowalski",
            "informant": "KA",
            "place": "Ocieszyn",
        },
        paragraph_speakers={0: "B", 1: "I"},
    )
    document.superscript_ranges = [(20, 21)]
    return document


def test_placeholder_names() -> None:
    assert placeholder_key("Miejscowość") == "miejscowosc"
    assert placeholder_key(" Gmina/powiat ") == "gmina_powiat"
    assert placeholder_key("Łęg") == "leg"
    custom = new_custom_field("Numer taśmy")
    names = placeholder_names((*DEFAULT_FIELDS, custom))
    assert names["sygnatura"] == ("field", "signature")
    assert names["signature"] == ("field", "signature")
    assert names["badacz"] == ("field", "researcher")
    assert names["numer_tasmy"] == ("field", custom.key)
    assert names["transkrypcja"] == ("special", "transkrypcja")


def test_template_is_filled_and_layout_kept(tmp_path) -> None:
    """ACC-55: układ szablonu zostaje, placeholdery są wypełnione."""
    template = make_template(tmp_path)
    target = tmp_path / "wynik.docx"

    report = export_docx(
        make_document(), target, options=ExportOptions(template_path=str(template))
    )

    read = DocxDocument(str(target))
    section = read.sections[0]
    margin = section.left_margin
    assert margin is not None and abs(margin - Cm(3.5)) < Cm(0.01)
    header = section.header
    assert header.paragraphs[0].text == " Sygnatura: AdK_1954"
    assert header.paragraphs[0]._p.xpath(".//w:drawing")  # logo zostało
    assert section.footer.paragraphs[0].text == "Badacz: J. Kowalski · {{nieznane}}"
    texts = [p.text for p in read.paragraphs]
    assert texts == [
        "Instytut Języka Polskiego",
        "B:\tRok urodzenia?",
        "I:\tTysionc dziewińćset",
        "Koniec dokumentu",
    ]
    style = read.paragraphs[1].style
    assert style is not None and style.name == "Quote"
    assert len(read.tables) == 1  # bez dodatkowej tabeli metryczki
    assert read.tables[0].rows[0].cells[1].text == "KA (Ocieszyn)"
    superscript = [
        r.text for p in read.paragraphs for r in p.runs if r.font.superscript
    ]
    assert superscript == ["n"]
    assert read.core_properties.title == "AdK_1954"
    assert any("{{nieznane}}" in warning for warning in report.warnings)


def test_anonymized_template_empties_personal_fields(tmp_path) -> None:
    template = make_template(tmp_path)
    target = tmp_path / "anonim.docx"
    export_docx(
        make_document(),
        target,
        anonymize=True,
        options=ExportOptions(template_path=str(template)),
    )
    read = DocxDocument(str(target))
    assert read.tables[0].rows[0].cells[1].text == " (Ocieszyn)"
    assert "KA" not in "".join(p.text for p in read.paragraphs)


def test_template_without_transcript_placeholder_appends_text(tmp_path) -> None:
    template = make_template(tmp_path, transcript=False)
    target = tmp_path / "koniec.docx"
    export_docx(
        make_document(), target, options=ExportOptions(template_path=str(template))
    )
    texts = [p.text for p in DocxDocument(str(target)).paragraphs]
    assert texts[-3:] == [
        "Koniec dokumentu",
        "B:\tRok urodzenia?",
        "I:\tTysionc dziewińćset",
    ]


def test_plain_template_gets_default_header_and_metadata_placeholder(tmp_path) -> None:
    plain = DocxDocument()
    plain.add_paragraph("Nagłówek instytucji")
    plain.add_paragraph("{{transkrypcja}}")
    plain_path = tmp_path / "prosty.docx"
    plain.save(str(plain_path))
    target = tmp_path / "a.docx"
    export_docx(
        make_document(), target, options=ExportOptions(template_path=str(plain_path))
    )
    read = DocxDocument(str(target))
    assert read.tables[0].rows[0].cells[1].text == "AdK_1954"
    # Import „najlepszego wysiłku” rozpoznaje mówców także w takim dokumencie.
    imported = import_docx(target)
    assert imported.text.endswith("Rok urodzenia?\nTysionc dziewińćset")
    assert sorted(imported.paragraph_speakers.values()) == ["B", "I"]

    with_table = DocxDocument()
    with_table.add_paragraph("{{metryczka}}")
    with_table.add_paragraph("Tekst:")
    table_path = tmp_path / "metryczka.docx"
    with_table.save(str(table_path))
    export_docx(
        make_document(), target, options=ExportOptions(template_path=str(table_path))
    )
    read = DocxDocument(str(target))
    assert [p.text for p in read.paragraphs][0] == "Tekst:"
    assert read.element.body[0].tag.endswith("tbl")


def test_check_template(tmp_path) -> None:
    template = make_template(tmp_path)
    check = check_template(template)
    assert check.ok
    assert check.unknown == ("nieznane",)
    assert "sygnatura" in check.placeholders

    broken = tmp_path / "zepsuty.docx"
    broken.write_bytes(b"to nie jest docx")
    assert not check_template(broken).ok
    assert not check_template(tmp_path / "brak.docx").ok


def test_broken_template_falls_back_to_builtin(tmp_path) -> None:
    """ACC-56: eksport kończy się sukcesem na szablonie wbudowanym."""
    broken = tmp_path / "zepsuty.docx"
    broken.write_bytes(b"to nie jest docx")
    target = tmp_path / "wynik.docx"

    report = export_docx(
        make_document(), target, options=ExportOptions(template_path=str(broken))
    )

    assert target.is_file()
    assert report.warnings and "zepsuty.docx" in report.warnings[0]
    assert DocxDocument(str(target)).tables[0].rows[0].cells[1].text == "AdK_1954"


def test_builtin_template_uses_export_font(tmp_path) -> None:
    target = tmp_path / "czcionka.docx"
    export_docx(
        make_document(), target, options=ExportOptions(font_family="Charis SIL")
    )
    assert DocxDocument(str(target)).styles["Normal"].font.name == "Charis SIL"
