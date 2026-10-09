"""Testy integracyjne eksportu DOCX z testowym odczytem wygenerowanych plików."""

import pytest
from docx import Document as DocxReader

from transkryptor.document.model import Document
from transkryptor.errors import ExportError
from transkryptor.export.docx_export import export_docx


class TestExport:
    """ACC-09: DOCX zawiera tekst, metadane i format indeksu górnego."""

    def test_docx_contains_text_metadata_and_superscript(self, tmp_path) -> None:
        doc = Document(
            text="pamientam uóna zmar-zły", author="Jan Kowalski", date="2026-09-30"
        )
        doc.superscript_ranges = [(5, 6)]  # nosowe „n” w „pamientam”
        target = tmp_path / "transkrypcja.docx"

        export_docx(doc, target)

        assert target.is_file()
        read = DocxReader(str(target))
        assert read.core_properties.author == "Jan Kowalski"

        metadata_paragraph = read.paragraphs[0]
        assert "Jan Kowalski" in metadata_paragraph.text
        assert "2026-09-30" in metadata_paragraph.text

        text_paragraph = read.paragraphs[1]
        assert text_paragraph.text == "pamientam uóna zmar-zły"
        superscript_runs = [r for r in text_paragraph.runs if r.font.superscript]
        assert len(superscript_runs) == 1
        assert superscript_runs[0].text == "n"
        plain_runs = [r for r in text_paragraph.runs if not r.font.superscript]
        # Pozostały tekst bez formatowania: całość bez nosowego „n”.
        assert "".join(r.text for r in plain_runs) == "pamietam uóna zmar-zły"

    def test_empty_document_exports(self, tmp_path) -> None:
        doc = Document()
        target = tmp_path / "pusty.docx"
        export_docx(doc, target)
        assert target.is_file()

    def test_failed_save_raises_export_error(self, tmp_path) -> None:
        doc = Document(text="tekst")
        with pytest.raises(ExportError) as exc_info:
            export_docx(doc, tmp_path / "nieistniejący" / "plik.docx")
        assert exc_info.value.user_message
        assert exc_info.value.retry_hint

    def test_polish_characters_preserved(self, tmp_path) -> None:
        doc = Document(
            text="żółć gęślą jaźń ąęćłńóśźż", author="Łukasz Żółw", date="1979-01-01"
        )
        target = tmp_path / "znaki.docx"
        export_docx(doc, target)
        read = DocxReader(str(target))
        assert "żółć gęślą jaźń" in read.paragraphs[1].text
        assert "Łukasz Żółw" in read.paragraphs[0].text


class TestMetadataTable:
    """ACC-32: tabela metryczki i eksport anonimizowany."""

    METADATA = {
        "signature": "AdK_1954",
        "place": "Ocieszyn",
        "informant": "KA",
        "informant_birth_year": "1954",
        "notes": "",
    }

    def _export(self, tmp_path, anonymize: bool):
        doc = Document(
            text="pamientam", author="Anna", date="2026-10-09", metadata=self.METADATA
        )
        doc.superscript_ranges = [(5, 6)]
        target = tmp_path / ("anonim.docx" if anonymize else "pelny.docx")
        export_docx(doc, target, anonymize=anonymize)
        return DocxReader(str(target))

    def test_table_contains_only_non_empty_fields(self, tmp_path) -> None:
        read = self._export(tmp_path, anonymize=False)
        rows = [[c.text for c in row.cells] for row in read.tables[0].rows]
        assert rows == [
            ["Sygnatura", "AdK_1954"],
            ["Miejscowość", "Ocieszyn"],
            ["Informator (kod)", "KA"],
            ["Rok urodzenia informatora", "1954"],
        ]
        assert read.core_properties.title == "AdK_1954"
        assert read.core_properties.keywords == "Ocieszyn"

    def test_anonymized_export_drops_personal_fields_only(self, tmp_path) -> None:
        read = self._export(tmp_path, anonymize=True)
        labels = [row.cells[0].text for row in read.tables[0].rows]
        assert labels == ["Sygnatura", "Miejscowość"]
        assert "KA" not in "".join(p.text for p in read.paragraphs)
        text = read.paragraphs[1]
        assert text.text == "pamientam"
        assert [r.text for r in text.runs if r.font.superscript] == ["n"]

    def test_without_metadata_there_is_no_table(self, tmp_path) -> None:
        target = tmp_path / "bez.docx"
        export_docx(Document(text="x"), target)
        read = DocxReader(str(target))
        assert read.tables == []
        assert read.core_properties.title == "Transkrypcja fonetyczna"
