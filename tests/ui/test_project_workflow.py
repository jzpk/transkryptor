"""Testy UI fazy 08: projekt, autozapis, metryczka i import DOCX (ACC-28…ACC-33)."""

import shutil
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document as DocxDocument
from PySide6.QtCore import QCoreApplication
from PySide6.QtGui import QTextCharFormat, QTextCursor
from PySide6.QtWidgets import QFileDialog, QMessageBox

from transkryptor.asr.engine import TranscriptionResult
from transkryptor.asr.manager import ModelManager
from transkryptor.asr.models import REQUIRED_FILES
from transkryptor.document.metadata import DEFAULT_FIELDS, encode_fields
from transkryptor.document.model import Document
from transkryptor.document.project import PlayerState, ProjectState, save_project
from transkryptor.export.docx_export import export_docx
from transkryptor.i18n import tr
from transkryptor.settings import MetadataSettings, ProjectSettings, Settings
from transkryptor.ui.asr_panel import APPLIED_MARK, PENDING_MARK
from transkryptor.ui.autosave import RecoveryDialog
from transkryptor.ui.main_window import DRAFT_REPLACE, MainWindow
from transkryptor.ui.settings_dialog import SettingsDialog

SAMPLE_MP3 = Path("tests/fixtures/audio/sample.mp3")
Button = QMessageBox.StandardButton


@pytest.fixture
def window(add_window) -> MainWindow:
    return add_window(MainWindow())


@pytest.fixture
def audio(tmp_path) -> Path:
    if not SAMPLE_MP3.is_file():
        pytest.skip(f"brak nagrania testowego {SAMPLE_MP3}")
    target = tmp_path / "nagrania" / "AdK_1954.mp3"
    target.parent.mkdir()
    shutil.copy(SAMPLE_MP3, target)
    return target


def save_dialog(monkeypatch, path: Path) -> None:
    monkeypatch.setattr(
        QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(path), ""))
    )


def open_dialog(monkeypatch, path: Path | None) -> None:
    result = (str(path), "") if path is not None else ("", "")
    monkeypatch.setattr(
        QFileDialog, "getOpenFileName", staticmethod(lambda *a, **k: result)
    )


def answer_questions(monkeypatch, answers: dict[str, Button]) -> list[tuple]:
    """``QMessageBox.question`` odpowiada wg tytułu; zwraca (tytuł, przyciski)."""
    calls: list[tuple] = []

    def fake(_parent, title, _text, buttons=Button.NoButton, *_args, **_kwargs):
        calls.append((title, buttons))
        return answers.get(title, Button.Cancel)

    monkeypatch.setattr(QMessageBox, "question", staticmethod(fake))
    return calls


def load_audio(qtbot, window: MainWindow, path: Path) -> None:
    window._on_import_audio(str(path))
    qtbot.waitUntil(lambda: window.player_bar.has_media, timeout=10000)


def fake_asr(window: MainWindow, tmp_path: Path) -> list[int]:
    """Imitacja modelu i transkrypcji; zwraca licznik wywołań ASR."""

    def downloader(repo_id, dest, on_progress, should_cancel):
        for name in REQUIRED_FILES:
            (dest / name).write_bytes(b"fake")

    calls: list[int] = []

    def transcribe(audio_path, model_dir, should_cancel, on_segment, on_progress=None):
        calls.append(1)
        segments = (
            SimpleNamespace(start_s=0.0, end_s=2.0, text="od", confidence=0.9),
            SimpleNamespace(start_s=2.0, end_s=3.5, text=" są", confidence=0.4),
        )
        for segment in segments:
            on_segment(segment)
        return TranscriptionResult(text="od są", language="pl", segments=segments)

    manager = ModelManager(models_root=tmp_path / "models", downloader=downloader)
    manager.download()
    window.asr_panel._manager = manager
    window.asr_panel._transcribe_impl = transcribe
    window.asr_panel.refresh_model_status()
    return calls


def run_asr(qtbot, window: MainWindow) -> None:
    window.asr_panel.transcribe_button.click()
    qtbot.waitUntil(lambda: window.asr_panel._transcribe_thread is None, timeout=10000)
    QCoreApplication.processEvents()


def review_labels(window: MainWindow) -> list[str]:
    items = window.asr_panel.suggestions_list
    return [items.item(i).text() for i in range(items.count())]


def superscripts(window: MainWindow) -> list[str]:
    window.run_validation_now()
    text = window.document.text
    return [text[s:e] for s, e in window.document.superscript_ranges]


class TestSaveAndOpen:
    """ACC-28: projekt odtwarza tekst, metryczkę, odtwarzacz, ASR i przegląd."""

    def test_round_trip_restores_work(
        self, qtbot, add_window, audio, tmp_path, monkeypatch
    ) -> None:
        window = add_window(MainWindow())
        load_audio(qtbot, window, audio)
        assert window.metadata_form.value("signature") == "AdK_1954"
        asr_calls = fake_asr(window, tmp_path)
        window.asr_panel.rule_checkboxes["SUG-LAB-U-INIT"].setChecked(False)
        run_asr(qtbot, window)
        assert window.editor.toPlainText() == "od som"
        cursor = window.editor.textCursor()
        cursor.movePosition(QTextCursor.MoveOperation.End)
        cursor.insertText(" koniec", QTextCharFormat())
        window.author_edit.setText("Anna Nowak")
        window.metadata_form.set_value("place", "Ocieszyn")
        window.player_bar.restore_state(
            PlayerState(position_ms=5000, loop_a_ms=1000, loop_b_ms=3000)
        )
        qtbot.waitUntil(lambda: window.player.position_ms == 5000, timeout=5000)

        path = tmp_path / "projekty" / "AdK.transkr"
        path.parent.mkdir()
        save_dialog(monkeypatch, path)
        window.save_action.trigger()
        assert path.is_file()
        assert not window.document.is_unsaved
        assert window.windowTitle() == "Transkryptor — AdK.transkr"
        assert window.settings_store.recent_projects()[0] == str(path)

        reopened = add_window(MainWindow())
        assert reopened.session.open_project(path)

        assert reopened.editor.toPlainText() == "od som koniec"
        assert superscripts(reopened) == ["m"]
        assert reopened.author_edit.text() == "Anna Nowak"
        assert reopened.metadata_form.value("signature") == "AdK_1954"
        assert reopened.metadata_form.value("place") == "Ocieszyn"
        assert not reopened.document.is_unsaved
        assert reopened.document.is_dirty  # zapisany, ale niewyeksportowany
        labels = review_labels(reopened)
        assert labels[0].startswith(PENDING_MARK) and "od" in labels[0]
        assert labels[1].startswith(APPLIED_MARK)
        assert len(reopened.editor.extraSelections()) == 1
        assert reopened.asr_panel.segments_list.count() == 2
        qtbot.waitUntil(lambda: reopened.player_bar.has_media, timeout=10000)
        qtbot.waitUntil(lambda: reopened.player.position_ms == 5000, timeout=5000)
        assert reopened.player_bar.state().loop_a_ms == 1000
        assert reopened.player_bar.state().loop_b_ms == 3000

        # Nieukończony przegląd: propozycję można zastosować po otwarciu.
        reopened._on_apply_review_item(0)
        assert reopened.editor.toPlainText() == "uod som koniec"
        # Ponowne wstawienie szkicu bez powtarzania ASR (REQ-14).
        reopened._ask_draft_placement = lambda: DRAFT_REPLACE
        reopened.asr_panel.reinsert_draft_button.click()
        # Reguły wybrane w panelu nowego okna (domyślnie wszystkie).
        assert reopened.editor.toPlainText() == "uod som"
        assert len(asr_calls) == 1

    def test_project_without_audio_opens_for_text_editing(
        self, qtbot, window, tmp_path, monkeypatch
    ) -> None:
        window.editor.setPlainText("sama edycja tekstu")
        path = tmp_path / "tekst.transkr"
        save_dialog(monkeypatch, path)
        assert window.session.save_project()
        window.editor.setPlainText("zmieniony")
        answer_questions(monkeypatch, {tr("session.unsaved.title"): Button.Discard})
        open_dialog(monkeypatch, path)
        window.open_action.trigger()
        assert window.editor.toPlainText() == "sama edycja tekstu"
        assert window.player.source_path is None
        assert not window.asr_panel.transcribe_button.isEnabled()

    def test_broken_project_shows_error_and_keeps_work(
        self, window, tmp_path, no_blocking_dialogs
    ) -> None:
        path = tmp_path / "zepsuty.transkr"
        path.write_text("{", encoding="utf-8")
        window.editor.setPlainText("praca")
        window.document.mark_exported()
        assert not window.session.open_project(path)
        assert window.editor.toPlainText() == "praca"
        assert any(kind == "warning" for kind, *_ in no_blocking_dialogs)

    def test_missing_recent_project_is_removed_from_list(
        self, window, tmp_path
    ) -> None:
        missing = str(tmp_path / "zniknął.transkr")
        window.settings_store.add_recent_project(missing)
        window.main_toolbar.refresh_recent_menu()
        names = [a.text() for a in window.main_toolbar.open_menu.actions()]
        assert "zniknął.transkr" in names
        assert not window.session.open_project(missing)
        assert window.settings_store.recent_projects() == []


class TestMovedOrChangedAudio:
    """ACC-29: przeniesione nagranie — prośba o plik; podmienione — ostrzeżenie."""

    def _project(self, qtbot, window, audio, tmp_path, monkeypatch) -> Path:
        load_audio(qtbot, window, audio)
        window.editor.setPlainText("tekst projektu")
        path = tmp_path / "p.transkr"
        save_dialog(monkeypatch, path)
        assert window.session.save_project()
        return path

    def test_moved_audio_can_be_located(
        self, qtbot, add_window, audio, tmp_path, monkeypatch
    ) -> None:
        path = self._project(
            qtbot, add_window(MainWindow()), audio, tmp_path, monkeypatch
        )
        moved = tmp_path / "gdzie-indziej.mp3"
        audio.rename(moved)
        calls = answer_questions(monkeypatch, {"Nie znaleziono nagrania": Button.Yes})
        open_dialog(monkeypatch, moved)

        other = add_window(MainWindow())
        assert other.session.open_project(path)

        assert [title for title, _ in calls] == ["Nie znaleziono nagrania"]
        assert other.player.source_path == moved
        assert other.editor.toPlainText() == "tekst projektu"

    def test_moved_audio_skipped_keeps_text(
        self, qtbot, add_window, audio, tmp_path, monkeypatch
    ) -> None:
        path = self._project(
            qtbot, add_window(MainWindow()), audio, tmp_path, monkeypatch
        )
        audio.unlink()
        answer_questions(monkeypatch, {"Nie znaleziono nagrania": Button.No})
        other = add_window(MainWindow())
        assert other.session.open_project(path)
        assert other.player.source_path is None
        assert other.editor.toPlainText() == "tekst projektu"

    @pytest.mark.parametrize("use_anyway", [True, False])
    def test_changed_audio_warns(
        self, qtbot, add_window, audio, tmp_path, monkeypatch, use_anyway
    ) -> None:
        path = self._project(
            qtbot, add_window(MainWindow()), audio, tmp_path, monkeypatch
        )
        with open(audio, "ab") as handle:
            handle.write(b"\0" * 16)  # inna suma SHA-256
        calls = answer_questions(
            monkeypatch, {"Inne nagranie": Button.Yes if use_anyway else Button.No}
        )
        other = add_window(MainWindow())
        assert other.session.open_project(path)
        assert [title for title, _ in calls] == ["Inne nagranie"]
        assert (other.player.source_path is not None) == use_anyway
        assert other.editor.toPlainText() == "tekst projektu"


class TestCloseQuestions:
    """ACC-31: pytanie zależy od tego, czy dokument ma plik projektu."""

    def test_saved_but_not_exported_project_closes_without_question(
        self, window, tmp_path, monkeypatch
    ) -> None:
        window.editor.setPlainText("zapisana praca")
        save_dialog(monkeypatch, tmp_path / "p.transkr")
        assert window.session.save_project()
        assert window.document.is_dirty  # niewyeksportowany
        calls = answer_questions(monkeypatch, {})
        assert window.session.can_close()
        assert calls == []

    def test_unsaved_changes_ask_with_save_option(
        self, window, tmp_path, monkeypatch
    ) -> None:
        path = tmp_path / "p.transkr"
        window.editor.setPlainText("wersja 1")
        save_dialog(monkeypatch, path)
        assert window.session.save_project()
        window.editor.setPlainText("wersja 2")
        assert window.windowTitle().endswith("p.transkr*")

        calls = answer_questions(
            monkeypatch, {tr("session.unsaved.title"): Button.Cancel}
        )
        assert not window.session.can_close()
        title, buttons = calls[0]
        assert title == tr("session.unsaved.title")
        assert buttons & Button.Save and buttons & Button.Discard

        answer_questions(monkeypatch, {tr("session.unsaved.title"): Button.Save})
        assert window.session.can_close()
        assert "wersja 2" in path.read_text(encoding="utf-8")

    def test_document_without_project_keeps_export_warning(
        self, window, tmp_path, monkeypatch
    ) -> None:
        window.editor.setPlainText("bez projektu")
        calls = answer_questions(
            monkeypatch, {tr("session.discard.title"): Button.Save}
        )
        target = tmp_path / "z-pytania.transkr"
        save_dialog(monkeypatch, target)
        assert window.session.can_close()
        assert calls[0][0] == tr("session.discard.title")
        assert target.is_file()  # „Zapisz” w oknie pytania zapisało projekt


class TestAutosaveAndRecovery:
    """ACC-30: autozapis, odzyskiwanie po awarii i sprzątanie po zamknięciu."""

    def test_autosave_writes_copy_without_touching_project(
        self, window, isolated_autosave
    ) -> None:
        autosave = window.session.autosave
        assert autosave.is_enabled
        assert not autosave.save_now()  # nic do zapisania
        window.editor.setPlainText("praca w toku")
        assert autosave.save_now()
        assert autosave.path.parent == isolated_autosave
        assert autosave.path.is_file()
        assert window.document.is_unsaved  # autozapis to nie zapis projektu
        assert window.session.project_path is None
        assert not autosave.save_now()  # bez zmian od ostatniego autozapisu

    def test_autosave_waits_for_asr(self, window) -> None:
        window.editor.setPlainText("tekst")
        window._on_transcription_started()
        try:
            assert not window.session.autosave.save_now()
            assert not window.session.autosave.path.exists()
        finally:
            window._on_transcription_finished()
        assert window.session.autosave.path.is_file()

    def test_disabled_autosave_removes_copy(self, window) -> None:
        window.editor.setPlainText("tekst")
        assert window.session.autosave.save_now()
        window.settings_store.save(
            Settings(project=ProjectSettings(autosave_enabled=False))
        )
        assert not window.session.autosave.path.exists()
        window.editor.setPlainText("dalej")
        assert not window.session.autosave.save_now()

    def test_clean_close_and_new_document_remove_copy(self, window) -> None:
        window.editor.setPlainText("tekst")
        autosave = window.session.autosave
        assert autosave.save_now()
        window.document.mark_exported()
        window._on_new_document()
        assert not autosave.path.exists()
        window.editor.setPlainText("drugi")
        assert autosave.save_now()
        window.document.mark_exported()
        assert window.session.can_close()
        assert not autosave.path.exists()

    def test_running_session_is_not_offered(self, window, add_window) -> None:
        window.editor.setPlainText("tekst")
        assert window.session.autosave.save_now()
        other = add_window(MainWindow())
        assert not other.session.offer_recovery()

    def _orphan(self, directory: Path, name: str, text: str) -> Path:
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / f"{name}.transkr"
        save_project(ProjectState(text=text, metadata={"signature": "AdK_1954"}), path)
        return path

    def test_recover_opens_unsaved_project(
        self, window, isolated_autosave, monkeypatch
    ) -> None:
        orphan = self._orphan(isolated_autosave, "dead", "odzyskany tekst jeż…")
        shown: list[list[str]] = []

        def fake_exec(dialog: RecoveryDialog) -> int:
            shown.append([c.label for c in dialog.candidates])
            dialog.recover_selected()
            return 1

        monkeypatch.setattr(RecoveryDialog, "exec", fake_exec)
        assert window.session.offer_recovery()

        assert "3 słowa" in shown[0][0] and "bez nagrania" in shown[0][0]
        assert window.editor.toPlainText() == "odzyskany tekst jeż…"
        assert window.metadata_form.value("signature") == "AdK_1954"
        assert window.session.project_path is None
        assert window.document.is_unsaved and window.document.is_dirty
        assert not orphan.exists()
        assert window.session.autosave.path.is_file()

    def test_discard_removes_orphan(
        self, window, isolated_autosave, monkeypatch
    ) -> None:
        first = self._orphan(isolated_autosave, "a", "pierwszy")
        second = self._orphan(isolated_autosave, "b", "drugi")

        def fake_exec(dialog: RecoveryDialog) -> int:
            dialog.discard_selected()
            return 0

        monkeypatch.setattr(RecoveryDialog, "exec", fake_exec)
        assert not window.session.offer_recovery()
        assert window.editor.toPlainText() == ""
        assert [first.exists(), second.exists()].count(True) == 1
        # Pozostały plik wraca przy następnym starcie.
        monkeypatch.setattr(RecoveryDialog, "exec", lambda dialog: 0)
        assert not window.session.offer_recovery()
        assert [first.exists(), second.exists()].count(True) == 1


class TestMetadataForm:
    def test_form_is_locked_during_asr(self, window) -> None:
        edit = window.metadata_form.edits["signature"]
        window._on_transcription_started()
        try:
            assert not edit.isEnabled()
        finally:
            window._on_transcription_finished()
        assert edit.isEnabled()

    def test_changes_reach_document_and_title(self, window) -> None:
        window.metadata_form.set_value("place", "Ocieszyn")
        assert window.document.metadata == {"place": "Ocieszyn"}
        assert window.windowTitle().endswith("*")
        assert "(1 z 8)" in window.metadata_form.toggle.text()

    def test_signature_hint_does_not_overwrite_or_dirty(
        self, qtbot, window, audio, tmp_path
    ) -> None:
        load_audio(qtbot, window, audio)
        assert window.metadata_form.value("signature") == "AdK_1954"
        assert not window.document.is_dirty  # sama podpowiedź to nie praca
        window.metadata_form.set_value("signature", "Moja_1")
        other = tmp_path / "Inne.mp3"
        shutil.copy(audio, other)
        window.document.mark_exported()
        window._on_import_audio(str(other))
        assert window.metadata_form.value("signature") == "Moja_1"

    def test_fields_follow_settings(self, window) -> None:
        dialog = SettingsDialog(window.settings_store.current, window)
        dialog.add_custom_field("Numer taśmy", personal=True)
        settings = dialog.settings()
        assert settings.metadata.fields[-1].label == "Numer taśmy"
        assert settings.metadata.fields[-1].personal
        window.settings_store.save(settings)
        assert len(window.metadata_form.edits) == 9

    def test_anonymized_export_from_toolbar(
        self, window, tmp_path, monkeypatch
    ) -> None:
        window.editor.setPlainText("tekst")
        window.metadata_form.set_value("signature", "AdK_1954")
        window.metadata_form.set_value("informant", "KA")
        target = tmp_path / "anonim.docx"
        save_dialog(monkeypatch, target)
        window.export_anonymized_action.trigger()
        read = DocxDocument(str(target))
        assert [row.cells[0].text for row in read.tables[0].rows] == ["Sygnatura"]
        assert not window.document.is_dirty


class TestSettingsSections:
    def test_project_and_metadata_sections_round_trip(self, window) -> None:
        reordered = (DEFAULT_FIELDS[2], *DEFAULT_FIELDS[:2], *DEFAULT_FIELDS[3:])
        custom = Settings(
            project=ProjectSettings(autosave_enabled=False, autosave_interval_s=300),
            metadata=MetadataSettings(fields_spec=encode_fields(reordered)),
        )
        dialog = SettingsDialog(custom, window)
        assert not dialog.autosave_spin.isEnabled()
        assert dialog.settings() == custom

    def test_reorder_and_remove_only_custom(self, window) -> None:
        dialog = SettingsDialog(Settings(), window)
        dialog.metadata_table.setCurrentCell(1, 0)
        dialog.field_up_button.click()
        assert dialog.metadata_fields()[0].key == "place"
        assert not dialog.remove_field_button.isEnabled()
        dialog.add_custom_field("Taśma")
        assert dialog.remove_field_button.isEnabled()
        dialog.remove_field_button.click()
        assert len(dialog.metadata_fields()) == len(DEFAULT_FIELDS)
        dialog.restore_defaults_button.click()
        assert dialog.settings() == Settings()


class TestImportDocx:
    """ACC-33 w UI: nowy, niezapisany projekt bez nagrania i raport."""

    def test_import_creates_unsaved_project(
        self, window, tmp_path, no_blocking_dialogs, monkeypatch
    ) -> None:
        source = Document(
            text="pamientam", author="Anna", metadata={"signature": "AdK_1954"}
        )
        source.superscript_ranges = [(5, 6)]
        path = tmp_path / "AdK.docx"
        export_docx(source, path)
        open_dialog(monkeypatch, path)

        window.import_docx_action.trigger()

        assert window.editor.toPlainText() == "pamientam"
        assert superscripts(window) == ["n"]
        assert window.author_edit.text() == "Anna"
        assert window.metadata_form.value("signature") == "AdK_1954"
        assert window.session.project_path is None
        assert window.player.source_path is None
        assert window.document.is_dirty
        info = [c for c in no_blocking_dialogs if c[1] == "Import DOCX"]
        assert info and "bez pominięć" in info[0][2]

    def test_import_is_guarded_like_new_document(
        self, window, tmp_path, monkeypatch
    ) -> None:
        window.editor.setPlainText("niewyeksportowana praca")
        answer_questions(monkeypatch, {tr("session.discard.title"): Button.Cancel})
        path = tmp_path / "x.docx"
        export_docx(Document(text="x"), path)
        assert not window.session.import_docx(path)
        assert window.editor.toPlainText() == "niewyeksportowana praca"
