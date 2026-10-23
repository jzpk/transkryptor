"""Testy UI mówców akapitów (faza 12, ACC-51)."""

from docx import Document as DocxDocument
from PySide6.QtCore import QPoint
from PySide6.QtGui import QTextCursor
from PySide6.QtWidgets import QMenu, QTableWidgetItem

from transkryptor.asr.engine import SegmentResult, TranscriptionResult
from transkryptor.document.speakers import INFORMANT, Speaker
from transkryptor.export.docx_export import export_docx
from transkryptor.notation.validator import VAL_12
from transkryptor.ui.main_window import MainWindow
from transkryptor.ui.shortcuts import SPEAKER_CLEAR_ID, speaker_id
from transkryptor.ui.speakers_dialog import SpeakersDialog

TEXT = "Rok urodzenia?\nTysionc dziewińćset\nA brat?"


def make_window(add_window) -> MainWindow:
    window = add_window(MainWindow())
    window.editor.load_content(TEXT, [])
    window.session.mark_editor_changed()
    return window


def cell(dialog: SpeakersDialog, row: int) -> QTableWidgetItem:
    item = dialog.table.item(row, 0)
    assert item is not None
    return item


def put_cursor(window: MainWindow, paragraph: int) -> None:
    block = window.editor.document().findBlockByNumber(paragraph)
    window.editor.setTextCursor(QTextCursor(block))


def assign(window: MainWindow, paragraph: int, number: int) -> None:
    put_cursor(window, paragraph)
    window.speaker_actions[speaker_id(number)].trigger()


def test_shortcuts_assign_speakers_and_undo_in_one_step(add_window) -> None:
    window = make_window(add_window)
    assert window.shortcut_actions[speaker_id(1)].shortcut().toString() == "Alt+Shift+1"

    assign(window, 0, 1)
    assign(window, 1, 2)

    assert window.document.paragraph_speakers == {0: "B", 1: "I"}
    assert window.document.text == TEXT  # etykieta nie jest tekstem
    assert window.editor.speaker_gutter.isVisibleTo(window)
    window.editor.undo()
    window.session.mark_editor_changed()
    assert window.document.paragraph_speakers == {0: "B"}

    # Ponowne użycie skrótu i Alt+Shift+0 zdejmują przypisanie.
    assign(window, 0, 1)
    assert window.document.paragraph_speakers == {}
    assign(window, 2, 1)
    window.speaker_actions[SPEAKER_CLEAR_ID].trigger()
    assert window.document.paragraph_speakers == {}


def test_selection_assigns_all_paragraphs(add_window) -> None:
    window = make_window(add_window)
    cursor = window.editor.textCursor()
    cursor.setPosition(2)
    cursor.movePosition(QTextCursor.MoveOperation.End, QTextCursor.MoveMode.KeepAnchor)
    window.editor.setTextCursor(cursor)
    window.speaker_actions[speaker_id(2)].trigger()
    assert window.document.paragraph_speakers == {0: "I", 1: "I", 2: "I"}


def test_missing_speaker_number_shows_message(add_window) -> None:
    window = make_window(add_window)
    assign(window, 0, 3)
    assert window.document.paragraph_speakers == {}
    assert "3" in window.statusBar().currentMessage()


def test_project_round_trip_and_docx_labels(add_window, tmp_path) -> None:
    """ACC-51: akapity i mówcy przetrwają zapis projektu; DOCX ma etykiety."""
    window = make_window(add_window)
    window.speakers.apply_speakers(
        [*window.document.speakers, Speaker("I2", "AK", INFORMANT)]
    )
    assign(window, 0, 1)
    assign(window, 1, 2)
    assign(window, 2, 3)
    path = tmp_path / "wywiad.transkr"
    assert window.session._write_project(path)

    reopened = add_window(MainWindow())
    assert reopened.session.open_project(path)
    assert reopened.document.paragraph_speakers == {0: "B", 1: "I", 2: "I2"}
    assert reopened.editor.paragraph_speakers() == {0: "B", 1: "I", 2: "I2"}
    assert "I2 — AK" in reopened.metadata_form.speakers_label.text()

    target = tmp_path / "wywiad.docx"
    export_docx(reopened.document, target)
    texts = [p.text for p in DocxDocument(str(target)).paragraphs]
    assert texts[-3:] == [
        "B:\tRok urodzenia?",
        "I:\tTysionc dziewińćset",
        "I2:\tA brat?",
    ]


def test_validator_hints_and_informant_word_count(add_window) -> None:
    window = make_window(add_window)
    window.run_validation_now()
    assert not [w for w in window.notation.warnings if w.code == VAL_12]

    assign(window, 0, 1)
    assign(window, 1, 2)
    window._update_title()
    window.run_validation_now()
    hints = [w for w in window.notation.warnings if w.code == VAL_12]
    assert [window.document.text[w.start : w.end] for w in hints] == ["A"]
    assert "(informator: 2)" in window.word_count_label.text()
    assert not window.search_bar.informant_check.isHidden()


def test_search_filter_limits_matches_to_informants(add_window) -> None:
    window = make_window(add_window)
    window.editor.load_content("tak tak\ntak\ntak", [], {0: "B", 1: "I"})
    window.session.mark_editor_changed()
    window._update_title()
    window.search.open_find()
    window.search_bar.find_edit.setText("tak")
    window.search.refresh()
    assert len(window.search.matches) == 4
    window.search_bar.informant_check.setChecked(True)
    window.search.refresh()
    assert window.search.matches == [(8, 11)]


def test_renaming_and_removing_speakers_updates_paragraphs(add_window) -> None:
    window = make_window(add_window)
    assign(window, 0, 1)
    assign(window, 1, 2)
    dialog = SpeakersDialog(window.document.speakers, {"B", "I"}, window)
    cell(dialog, 0).setText("R")
    dialog.table.setCurrentCell(1, 0)
    dialog.remove_selected()
    assert dialog.renamed() == {"B": "R"}
    assert dialog.removed() == {"I"}
    window.speakers.apply_speakers(
        dialog.speakers(), dialog.renamed(), dialog.removed()
    )
    assert window.document.paragraph_speakers == {0: "R"}
    assert window.editor.paragraph_speakers() == {0: "R"}


def test_dialog_rejects_invalid_keys(add_window) -> None:
    window = make_window(add_window)
    dialog = SpeakersDialog(window.document.speakers, (), window)
    dialog.add_speaker()
    assert dialog.speakers()[-1].key == "I2"
    cell(dialog, 2).setText("B")
    assert "B" in (dialog.validation_error() or "")
    cell(dialog, 2).setText("a b")
    assert dialog.validation_error()


def test_context_menu_lists_speakers(add_window) -> None:
    window = make_window(add_window)
    menu = QMenu()
    window.speakers.extend_context_menu(menu, 0)
    submenu = menu.actions()[-1].menu()
    assert isinstance(submenu, QMenu)
    texts = [action.text() for action in submenu.actions() if action.text()]
    assert texts[0].startswith("B — Badacz")
    submenu.actions()[1].trigger()
    assert window.document.paragraph_speakers == {0: "I"}


def test_gutter_click_opens_speaker_menu(add_window, qtbot, monkeypatch) -> None:
    window = make_window(add_window)
    window.show()
    qtbot.waitExposed(window)
    assign(window, 0, 1)
    shown: list[int] = []
    monkeypatch.setattr(
        window.speakers,
        "exec_menu",
        lambda menu, _pos: shown.append(len(menu.actions())),
    )
    block = window.editor.speaker_gutter.block_at(5)
    assert block is not None and block.blockNumber() == 0
    window.editor.speaker_gutter._on_click(block, QPoint(0, 0))
    assert shown and shown[0] >= 3


def diarized_result(*turns: tuple[float, float, str, int]) -> TranscriptionResult:
    segments = tuple(
        SegmentResult(start, end, text, 0.9, speaker=speaker)
        for start, end, text, speaker in turns
    )
    return TranscriptionResult(
        text="".join(s.text for s in segments).strip(),
        language="pl",
        segments=segments,
    )


DIARIZED = diarized_result(
    (0.0, 2.0, " Gdzie pani mieszka?", 0),
    (2.0, 5.0, " Tu, we wsi.", 1),
    (5.0, 7.0, " A brat?", 0),
)


def test_diarized_draft_gets_speaker_paragraphs(add_window, tmp_path) -> None:
    """Podział na mówców: wypowiedź = akapit z mówcą, także w DOCX."""
    window = add_window(MainWindow())
    window._on_draft_ready(DIARIZED)
    window.session.mark_editor_changed()

    assert window.document.text == (
        "[00:00] Gdzie pani mieszka?\n[00:02] Tu, we wsi.\n[00:05] A brat?"
    )
    # Pierwszy głos to pierwszy mówca z listy (B), drugi — I.
    assert window.document.paragraph_speakers == {0: "B", 1: "I", 2: "B"}
    assert window.editor.speaker_gutter.isVisibleTo(window)

    target = tmp_path / "wywiad.docx"
    export_docx(window.document, target)
    texts = [p.text for p in DocxDocument(str(target)).paragraphs]
    assert texts[-3:] == [
        "B:\t[00:00] Gdzie pani mieszka?",
        "I:\t[00:02] Tu, we wsi.",
        "B:\t[00:05] A brat?",
    ]

    # Wstawienie razem z mówcami to jeden krok cofania.
    window.editor.undo()
    window.session.mark_editor_changed()
    assert window.document.text == ""
    assert window.document.paragraph_speakers == {}


def test_extra_voices_get_new_informants(add_window) -> None:
    window = add_window(MainWindow())
    window._on_draft_ready(
        diarized_result(
            (0.0, 1.0, " Pytanie.", 0),
            (1.0, 2.0, " Odpowiedź.", 1),
            (2.0, 3.0, " Wtrącenie.", 2),
        )
    )
    window.session.mark_editor_changed()
    assert [s.key for s in window.document.speakers] == ["B", "I", "I2"]
    assert window.document.speaker("I2") == Speaker("I2", "", INFORMANT)
    assert window.document.paragraph_speakers == {0: "B", 1: "I", 2: "I2"}


def test_cancelled_placement_keeps_speaker_list(add_window, monkeypatch) -> None:
    window = make_window(add_window)
    monkeypatch.setattr(window, "_ask_draft_placement", lambda: None)
    window._on_draft_ready(
        diarized_result((0.0, 1.0, " a", 0), (1.0, 2.0, " b", 1), (2.0, 3.0, " c", 2))
    )
    assert window.document.text == TEXT
    assert [s.key for s in window.document.speakers] == ["B", "I"]
