"""Integralność tekstu w oknie: edytor = projekt = DOCX (P1 audytu).

Znaki spoza BMP (BUG-02), niełamliwa spacja (BUG-03) i synchronizacja
z debounce (PERF-02).
"""

import pytest
from docx import Document as DocxDocument
from PySide6.QtGui import QTextCharFormat

from transkryptor.document.project import ReviewEntry, load_project
from transkryptor.export.docx_export import export_docx
from transkryptor.ui.main_window import MainWindow

SUPERSCRIPT = QTextCharFormat.VerticalAlignment.AlignSuperScript


@pytest.fixture
def window(add_window) -> MainWindow:
    return add_window(MainWindow())


def superscript_text(path) -> str:
    docx = DocxDocument(str(path))
    return "".join(
        run.text
        for paragraph in docx.paragraphs
        for run in paragraph.runs
        if run.font.superscript
    )


def test_emoji_before_superscript_reaches_project_and_docx(window, tmp_path) -> None:
    window.editor.load_content("😀 bendzie", [(4, 5)])
    shown = window.editor.track_range(4, 5)
    assert shown.selectedText() == "n"
    assert shown.charFormat().verticalAlignment() == SUPERSCRIPT
    project = tmp_path / "p.transkr"
    window.session._write_project(project)
    state = load_project(project)
    assert state.text == "😀 bendzie"
    assert list(state.superscript_ranges) == [(4, 5)]
    docx = tmp_path / "p.docx"
    export_docx(window.document, docx)
    assert superscript_text(docx) == "n"


def test_nbsp_round_trips_through_project_and_docx(window, tmp_path) -> None:
    window.editor.load_content("na pewno", [])
    window.editor.textCursor().insertText("x")
    project = tmp_path / "p.transkr"
    window.session._write_project(project)
    assert " " in load_project(project).text
    docx = tmp_path / "p.docx"
    export_docx(window.document, docx)
    assert " " in DocxDocument(str(docx)).paragraphs[-1].text


def test_save_right_after_typing_includes_change(window, tmp_path) -> None:
    window.editor.setPlainText("bendzie")
    window.editor.textCursor().insertText("x")
    project = tmp_path / "p.transkr"
    window.session._write_project(project)  # bez czekania na debounce
    assert load_project(project).text == "xbendzie"
    assert not window.document.is_unsaved


def test_typing_marks_title_immediately_and_syncs_later(qtbot, window) -> None:
    revision = window.session._document.revision
    window.editor.textCursor().insertText("a")
    assert window.windowTitle().endswith("*")
    assert window.session._document.revision == revision  # tylko znacznik
    qtbot.waitUntil(lambda: window.session._document.text == "a")


def test_superscript_toggle_syncs_once(qtbot, window) -> None:
    window.editor.setPlainText("bendzie")
    window.document.mark_saved()
    revision = window.document.revision
    window.editor.setTextCursor(window.editor.track_range(1, 2))
    window._on_superscript()
    assert window.document.superscript_ranges == [(1, 2)]
    assert window.document.revision == revision + 1


def test_warning_navigation_after_emoji(window) -> None:
    window.editor.setPlainText("😀 bendzie")
    window.notation.go_to_range(2, 9)
    assert window.editor.textCursor().selectedText() == "bendzie"


def test_review_entries_after_emoji_use_document_positions(window) -> None:
    window.editor.setPlainText("😀 bendzie")
    entry = ReviewEntry(
        start=2,
        end=9,
        applied=False,
        code="X",
        suggestion_start=2,
        suggestion_end=9,
        original="bendzie",
        replacement="będzie",
    )
    window.review.restore([entry])
    assert window.review.items[0].cursor.selectedText() == "bendzie"
    assert (window.review.entries()[0].start, window.review.entries()[0].end) == (2, 9)
