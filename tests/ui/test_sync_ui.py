"""Tekst powiązany z nagraniem w oknie (faza 11: ACC-45, ACC-46, ACC-48)."""

import shutil
import time
from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from PySide6.QtCore import Qt
from PySide6.QtGui import QTextCursor
from PySide6.QtWidgets import QFileDialog

from transkryptor.asr.engine import SegmentResult, TranscriptionResult, WordResult
from transkryptor.document.anchors import MANUAL, SEGMENT
from transkryptor.ui.highlights import ACTIVE_LAYER, UNCERTAIN_LAYER
from transkryptor.ui.main_window import MainWindow

SAMPLE_MP3 = Path("tests/fixtures/audio/sample.mp3")

SENTENCES = (
    (0.0, 2.0, " Pierwsze zdanie."),
    (2.0, 4.0, " Drugie zdanie."),
    (4.0, 6.5, " Trzecie zdanie."),
    (6.5, 9.0, " Czwarte zdanie."),
)


def draft_result(segments=SENTENCES) -> TranscriptionResult:
    results = tuple(SegmentResult(s, e, text, 0.9) for s, e, text in segments)
    return TranscriptionResult(
        text="".join(r.text for r in results).strip(), language="pl", segments=results
    )


@pytest.fixture
def audio(tmp_path) -> Path:
    if not SAMPLE_MP3.is_file():
        pytest.skip(f"brak nagrania testowego {SAMPLE_MP3}")
    target = tmp_path / "nagranie.mp3"
    shutil.copy(SAMPLE_MP3, target)
    return target


@pytest.fixture
def window(qtbot, add_window, audio) -> MainWindow:
    win = add_window(MainWindow())
    win.resize(1000, 700)
    win.show()
    win._on_import_audio(str(audio))
    qtbot.waitUntil(lambda: win.player_bar.has_media, timeout=10000)
    return win


@pytest.fixture
def playback(window, monkeypatch) -> list[tuple[str, int]]:
    """Zlecenia odtwarzacza zamiast prawdziwego odtwarzania."""
    calls: list[tuple[str, int]] = []
    player = window.player
    monkeypatch.setattr(player, "set_position", lambda ms: calls.append(("seek", ms)))
    monkeypatch.setattr(player, "play", lambda: calls.append(("play", 0)))
    monkeypatch.setattr(
        player,
        "play_loop_once",
        lambda start, end: calls.extend([("loop", start), ("to", end)]),
    )
    return calls


def ctrl_click(qtbot, window: MainWindow, position: int) -> None:
    cursor = QTextCursor(window.editor.document())
    cursor.setPosition(position)
    point = window.editor.cursorRect(cursor).center()
    qtbot.mouseClick(
        window.editor.viewport(),
        Qt.MouseButton.LeftButton,
        Qt.KeyboardModifier.ControlModifier,
        point,
    )


def segment_texts(window: MainWindow) -> list[str]:
    text = window.editor.plain_text()
    return [text[a.start : a.end] for a in window.sync.entries() if a.source == SEGMENT]


class TestPlaybackFromText:
    def test_acc45_ctrl_click_and_active_segment(self, qtbot, window, playback) -> None:
        window._on_draft_ready(draft_result())
        assert segment_texts(window) == [s[2].strip() for s in SENTENCES]
        text = window.editor.plain_text()
        third = text.index("Trzecie") + 3
        ctrl_click(qtbot, window, third)
        preroll = window.settings_store.current.player.segment_preroll_ms
        assert playback == [("seek", 4000 - preroll), ("play", 0)]

        window.player.position_changed.emit(500)
        layer = window.editor.highlight_layer(ACTIVE_LAYER)
        assert [s.cursor.selectedText() for s in layer] == ["Pierwsze zdanie."]
        qtbot.wait(200)  # przerwa dłuższa niż ograniczenie częstotliwości
        window.player.position_changed.emit(7000)
        qtbot.waitUntil(
            lambda: [
                s.cursor.selectedText()
                for s in window.editor.highlight_layer(ACTIVE_LAYER)
            ]
            == ["Czwarte zdanie."],
            timeout=2000,
        )

    def test_play_from_cursor_shortcut_uses_segment(self, window, playback) -> None:
        window._on_draft_ready(draft_result())
        cursor = window.editor.textCursor()
        cursor.setPosition(window.editor.plain_text().index("Drugie") + 2)
        window.editor.setTextCursor(cursor)
        window.player_actions["play_from_cursor"].trigger()
        preroll = window.settings_store.current.player.segment_preroll_ms
        assert playback[0] == ("seek", 2000 - preroll)

    def test_text_without_anchor_reports_status(self, qtbot, window, playback) -> None:
        window.editor.setPlainText("Tekst wpisany ręcznie.")
        ctrl_click(qtbot, window, 3)
        assert playback == []
        assert "nie jest powiązane" in window.statusBar().currentMessage()


class TestAnchorsSurviveEditing:
    def test_acc46_delete_segment_edit_another_and_reopen(
        self, qtbot, window, add_window, tmp_path, monkeypatch
    ) -> None:
        window._on_draft_ready(draft_result())
        editor = window.editor
        text = editor.plain_text()
        # Usunięcie całego drugiego segmentu (ze spacją przed nim).
        cursor = QTextCursor(editor.document())
        start = text.index(" Drugie")
        cursor.setPosition(start)
        cursor.setPosition(
            start + len(" Drugie zdanie."), QTextCursor.MoveMode.KeepAnchor
        )
        cursor.removeSelectedText()
        # Dopisanie tekstu w środku trzeciego segmentu.
        cursor.setPosition(editor.plain_text().index("zdanie", 20))
        cursor.insertText("bardzo ")
        expected = ["Pierwsze zdanie.", "Trzecie bardzo zdanie.", "Czwarte zdanie."]
        assert segment_texts(window) == expected

        path = tmp_path / "projekt.transkr"
        monkeypatch.setattr(
            QFileDialog,
            "getSaveFileName",
            staticmethod(lambda *a, **k: (str(path), "")),
        )
        assert window.session.save_project()
        reopened = add_window(MainWindow())
        assert reopened.session.open_project(path)
        assert segment_texts(reopened) == expected
        times = [
            (a.start_ms, a.end_ms)
            for a in reopened.sync.entries()
            if a.source == SEGMENT
        ]
        assert times == [(0, 2000), (4000, 6500), (6500, 9000)]

    def test_replace_draft_drops_old_anchors(self, window, monkeypatch) -> None:
        window._on_draft_ready(draft_result())
        monkeypatch.setattr(window, "_ask_draft_placement", lambda: "replace")
        window._on_draft_ready(draft_result(SENTENCES[:1]))
        assert segment_texts(window) == ["Pierwsze zdanie."]

    def test_append_draft_keeps_offsets(self, window, monkeypatch) -> None:
        window._on_draft_ready(draft_result(SENTENCES[:2]))
        monkeypatch.setattr(window, "_ask_draft_placement", lambda: "append")
        window._on_draft_ready(draft_result(SENTENCES[2:]))
        assert segment_texts(window) == [s[2].strip() for s in SENTENCES]

    def test_rule_at_segment_end_stays_in_segment(self, window) -> None:
        for code, box in window.asr_panel.rule_checkboxes.items():
            box.setChecked(code == "SUG-NAS-A-FINAL")
        window._on_draft_ready(
            draft_result(((0.0, 1.0, " oni są"), (1.0, 2.0, " i już")))
        )
        assert window.editor.plain_text() == "[00:00] oni som i już"
        assert segment_texts(window) == ["oni som", "i już"]
        cursor = QTextCursor(window.editor.document())
        cursor.setPosition(len("[00:00] oni som"))
        cursor.insertText(" no")
        assert segment_texts(window) == ["oni som", "i już"]


class TestManualTimestamps:
    def test_acc48_timestamp_plays_from_saved_time_and_is_not_exported(
        self, qtbot, window, tmp_path, monkeypatch
    ) -> None:
        editor = window.editor
        editor.setPlainText("Pierwszy akapit tekstu. Drugi fragment.")
        window.document.mark_exported()
        cursor = editor.textCursor()
        marker_at = editor.plain_text().index("Drugi")
        cursor.setPosition(marker_at)
        editor.setTextCursor(cursor)
        monkeypatch.setattr(
            type(window.player), "position_ms", property(lambda _: 12345)
        )
        window.player_actions["insert_timestamp"].trigger()
        assert [(a.start, a.start_ms, a.source) for a in window.sync.entries()] == [
            (marker_at, 12345, MANUAL)
        ]
        # Znacznik to dane projektu, nie tekst eksportu.
        assert window.document.is_unsaved
        assert not window.document.is_dirty

        calls: list[tuple[str, int]] = []
        monkeypatch.setattr(
            window.player, "set_position", lambda ms: calls.append(("seek", ms))
        )
        monkeypatch.setattr(window.player, "play", lambda: calls.append(("play", 0)))
        ctrl_click(qtbot, window, marker_at + 4)
        assert calls == [("seek", 12345), ("play", 0)]

        target = tmp_path / "eksport.docx"
        monkeypatch.setattr(
            QFileDialog,
            "getSaveFileName",
            staticmethod(lambda *a, **k: (str(target), "")),
        )
        window._on_export()
        paragraphs = [p.text for p in DocxDocument(str(target)).paragraphs]
        # Decyzja fazy 11: znaczniki niewidoczne w DOCX.
        assert paragraphs[-1] == "Pierwszy akapit tekstu. Drugi fragment."

    def test_timestamp_moves_with_text_and_can_be_removed(self, window) -> None:
        editor = window.editor
        editor.setPlainText("abc def")
        cursor = editor.textCursor()
        cursor.setPosition(4)
        editor.setTextCursor(cursor)
        window.sync.insert_timestamp()
        cursor.setPosition(0)
        cursor.insertText("xyz ")
        assert [a.start for a in window.sync.entries()] == [8]
        # Tekst wpisany w miejscu znacznika zaczyna się od znacznika.
        cursor.setPosition(8)
        cursor.insertText("nowe ")
        assert [a.start for a in window.sync.entries()] == [8]
        item = window.sync.manual_near(8)
        assert item is not None
        window.sync.remove(item)
        assert window.sync.entries() == []


def words_result() -> TranscriptionResult:
    """Szkic z atrapy ASR: słowa z pewnością poniżej i powyżej progu (0,4)."""
    plan = [
        (0.0, 2.0, [("Wtedy", 0.95), ("my", 0.30), ("żeśmy", 0.12)]),
        (2.0, 4.0, [("poszli", 0.85), ("do", 0.39), ("lasu.", 0.41)]),
    ]
    segments = []
    for start, end, words in plan:
        step = (end - start) / len(words)
        results = tuple(
            WordResult(start + i * step, start + (i + 1) * step, f" {w}", p)
            for i, (w, p) in enumerate(words)
        )
        text = "".join(w.text for w in results)
        segments.append(SegmentResult(start, end, text, 0.7, results))
    text = "".join(s.text for s in segments).strip()
    return TranscriptionResult(text=text, language="pl", segments=tuple(segments))


class TestUncertainWords:
    def test_acc47_uncertain_words(
        self, qtbot, window, playback, tmp_path, monkeypatch
    ) -> None:
        window._on_draft_ready(words_result())
        sync = window.sync
        assert sync.uncertain_words() == ["my", "żeśmy", "do"]
        layer = window.editor.highlight_layer(UNCERTAIN_LAYER)
        assert [s.cursor.selectedText() for s in layer] == ["my", "żeśmy", "do"]

        # Ctrl+klik w słowo odtwarza dokładnie jego zakres (z marginesem).
        text = window.editor.plain_text()
        ctrl_click(qtbot, window, text.index("żeśmy") + 2)
        assert playback == [("loop", 4000 // 3 - 150), ("to", 2000 + 150)]

        # Poprawione słowo traci podkreślenie; reszta zostaje.
        cursor = QTextCursor(window.editor.document())
        cursor.setPosition(text.index("żeśmy"))
        cursor.setPosition(text.index("żeśmy") + 5, QTextCursor.MoveMode.KeepAnchor)
        cursor.insertText("żeźmy")
        sync.update_uncertain()
        assert sync.uncertain_words() == ["my", "do"]
        layer = window.editor.highlight_layer(UNCERTAIN_LAYER)
        assert [s.cursor.selectedText() for s in layer] == ["my", "do"]
        assert "pewność 30%" in (
            sync.tooltip_at(window.editor.plain_text().index("my")) or ""
        )

        # Poprawka przetrwa zapis projektu; podkreślenia nie trafiają do DOCX.
        entries = sync.entries()
        assert [a.edited for a in entries if a.source == "word"] == [
            False,
            False,
            True,
            False,
            False,
            False,
        ]
        target = tmp_path / "eksport.docx"
        monkeypatch.setattr(
            QFileDialog,
            "getSaveFileName",
            staticmethod(lambda *a, **k: (str(target), "")),
        )
        window._on_export()
        runs = [run for p in DocxDocument(str(target)).paragraphs for run in p.runs]
        assert runs and not any(run.underline for run in runs)

    def test_threshold_and_switch_from_settings(self, window) -> None:
        window._on_draft_ready(words_result())
        current = window.settings_store.current
        window.settings_store.save(
            replace(current, asr=replace(current.asr, word_threshold_pct=20))
        )
        assert window.sync.uncertain_words() == ["żeśmy"]
        window.settings_store.save(
            replace(current, asr=replace(current.asr, show_uncertain_words=False))
        )
        assert window.sync.uncertain_words() == []
        assert window.editor.highlight_layer(UNCERTAIN_LAYER) == []

    def test_words_survive_project_round_trip(
        self, window, add_window, tmp_path, monkeypatch
    ) -> None:
        window.asr_panel.restore_result(words_result())
        window._on_draft_ready(words_result())
        path = tmp_path / "projekt.transkr"
        monkeypatch.setattr(
            QFileDialog,
            "getSaveFileName",
            staticmethod(lambda *a, **k: (str(path), "")),
        )
        assert window.session.save_project()
        reopened = add_window(MainWindow())
        assert reopened.session.open_project(path)
        assert reopened.sync.uncertain_words() == ["my", "żeśmy", "do"]
        result = reopened.asr_panel.last_result
        assert result is not None and result.segments[0].words[1].probability == 0.3


def test_sixty_minute_document_stays_responsive(qtbot, add_window) -> None:
    """Wydajność: szkic odpowiadający ~60 min nagrania (≈9000 słów)."""
    window = add_window(MainWindow())
    window.resize(1000, 700)
    window.show()
    segments = []
    for index in range(1000):
        start = index * 3.6
        words = tuple(
            WordResult(
                start + i * 0.4, start + i * 0.4 + 0.4, f" słowo{i}", 0.3 + i / 20
            )
            for i in range(9)
        )
        segments.append(
            SegmentResult(
                start, start + 3.6, "".join(w.text for w in words), 0.6, words
            )
        )
    result = TranscriptionResult(
        text="".join(s.text for s in segments).strip(),
        language="pl",
        segments=tuple(segments),
    )
    began = time.perf_counter()
    window._on_draft_ready(result)
    insert_s = time.perf_counter() - began
    assert len(window.sync.items) == 10_000

    began = time.perf_counter()
    for _ in range(20):
        qtbot.keyClick(window.editor, Qt.Key.Key_A)
    key_ms = (time.perf_counter() - began) / 20 * 1000

    began = time.perf_counter()
    window.sync.update_uncertain()
    uncertain_ms = (time.perf_counter() - began) * 1000

    began = time.perf_counter()
    for position in range(0, 3_600_000, 36_000):
        window.player.position_changed.emit(position)
        window.sync._update_active()
    tick_ms = (time.perf_counter() - began) / 100 * 1000

    print(
        f"\nwstawienie {insert_s:.2f} s, klawisz {key_ms:.1f} ms, "
        f"niepewne {uncertain_ms:.1f} ms, podświetlenie {tick_ms:.2f} ms"
    )
    assert insert_s < 10
    assert key_ms < 50
    assert uncertain_ms < 100
    assert tick_ms < 20
