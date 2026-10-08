"""Test end-to-end całej ścieżki pracy (faza 07, propozycja 26).

Import nagrania → ASR (prawdziwy ``asr.engine.transcribe`` z atrapą backendu)
→ szkic w edytorze → automatyczne reguły → przegląd → wyszukaj i zamień
z indeksem górnym → ujednolicenie wielokropków → eksport DOCX → odczyt DOCX.
"""

from functools import partial
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document as DocxDocument
from PySide6.QtCore import QDate
from PySide6.QtGui import QTextCursor
from PySide6.QtWidgets import QFileDialog

from transkryptor.asr import engine
from transkryptor.asr.manager import ModelManager
from transkryptor.asr.models import REQUIRED_FILES
from transkryptor.ui.main_window import MainWindow

SAMPLE_MP3 = "test/JaE_1979_przesądy.mp3"

SEGMENTS = (
    (0.0, 2.0, "od tego czasu ... będzie"),
    (2.0, 4.0, " jeż... kam... są."),
    (4.0, 6.0, " A będzie lepiej."),
)


def fake_downloader(repo_id: str, dest: Path, on_progress, should_cancel) -> None:
    for name in REQUIRED_FILES:
        (dest / name).write_bytes(b"fake")
        on_progress(1024, 4096, name)


def fake_backend_factory(model_dir):
    """Imitacja faster-whisper: segmenty z wielokropkami ASCII, bez modelu."""

    class FakeBackend:
        def transcribe(self, audio_path, language, vad_filter):
            segments = [
                SimpleNamespace(start=start, end=end, text=text, avg_logprob=-0.1)
                for start, end, text in SEGMENTS
            ]
            return iter(segments), SimpleNamespace(language=language, duration=6.0)

    return FakeBackend()


@pytest.fixture
def window(qtbot, add_window, tmp_path):
    if not Path(SAMPLE_MP3).is_file():
        pytest.skip(f"brak nagrania testowego {SAMPLE_MP3}")
    win = add_window(MainWindow())
    manager = ModelManager(models_root=tmp_path / "models", downloader=fake_downloader)
    manager.download()
    win.asr_panel._manager = manager
    win.asr_panel._transcribe_impl = partial(
        engine.transcribe, backend_factory=fake_backend_factory
    )
    win.asr_panel.refresh_model_status()
    return win


def superscripts(window: MainWindow) -> list[str]:
    window.run_validation_now()
    text = window.document.text
    return [text[start:end] for start, end in window.document.superscript_ranges]


def test_full_workflow(qtbot, window, tmp_path, monkeypatch) -> None:
    # 1. Import nagrania.
    window._on_import_audio(SAMPLE_MP3)
    qtbot.waitUntil(lambda: window.player_bar.has_media, timeout=10000)

    # 2. ASR: automatycznie tylko nosowość w wygłosie (są → som).
    for code, box in window.asr_panel.rule_checkboxes.items():
        box.setChecked(code == "SUG-NAS-A-FINAL")
    window.asr_panel.transcribe_button.click()
    qtbot.waitUntil(
        lambda: window.asr_panel._transcribe_thread is None
        and window.editor.toPlainText() != "",
        timeout=10000,
    )
    assert window.asr_panel.segments_list.count() == len(SEGMENTS)

    # 3. Szkic w stylu wielokropka z ustawień (domyślnie „…”), reguła zastosowana.
    assert window.editor.toPlainText() == (
        "od tego czasu … będzie jeż… kam… som. A będzie lepiej."
    )
    assert superscripts(window) == ["m"]

    # 4. Przegląd: zastosowanie propozycji labializacji dla „od”.
    index = next(
        i
        for i, item in enumerate(window.review.items)
        if item.suggestion.code == "SUG-LAB-U-INIT" and item.suggestion.start == 0
    )
    window._on_apply_review_item(index)
    assert window.editor.toPlainText().startswith("uod tego")
    assert superscripts(window) == ["u", "m"]

    # 5. Wyszukaj i zamień wszystkie z indeksem górnym (jedno cofnięcie).
    window.search_actions["replace"].trigger()
    window.search_bar.whole_words_check.setChecked(True)
    window.search_bar.find_edit.setText("będzie")
    window.search_bar.replace_edit.setText("be^ndzie")
    window.search_bar.replace_all_button.click()
    assert window.editor.toPlainText() == (
        "uod tego czasu … bendzie jeż… kam… som. A bendzie lepiej."
    )
    assert superscripts(window) == ["u", "n", "m", "n"]
    window.search.close()

    # 6. Ręcznie dopisana pauza ASCII i „Ujednolić wielokropki”.
    cursor = window.editor.textCursor()
    cursor.movePosition(QTextCursor.MoveOperation.End)
    window.editor.setTextCursor(cursor)
    window.editor.insertPlainText(" I tak ... no.")
    window.run_validation_now()
    assert window.warnings_panel.warning_count() == 1  # VAL-05
    window.unify_ellipses_action.trigger()
    expected = "uod tego czasu … bendzie jeż… kam… som. A bendzie lepiej. I tak … no."
    assert window.editor.toPlainText() == expected
    assert window.warnings_panel.warning_count() == 0

    # 7. Metadane i eksport DOCX.
    window.author_edit.setText("Anna Nowak")
    window.date_edit.setDate(QDate(2026, 10, 8))
    path = tmp_path / "wynik.docx"
    monkeypatch.setattr(
        QFileDialog,
        "getSaveFileName",
        staticmethod(lambda *args, **kwargs: (str(path), "")),
    )
    window.export_action.trigger()
    assert not window.document.is_dirty

    # 8. Odczyt DOCX: tekst, indeksy górne i metadane.
    docx = DocxDocument(str(path))
    assert docx.core_properties.author == "Anna Nowak"
    paragraphs = [p for p in docx.paragraphs if "tego czasu" in p.text]
    assert [p.text for p in paragraphs] == [expected]
    runs = paragraphs[0].runs
    assert [run.text for run in runs if run.font.superscript] == ["u", "n", "m", "n"]
    assert any(
        "Anna Nowak" in p.text and "2026-10-08" in p.text for p in docx.paragraphs
    )
