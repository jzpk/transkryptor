"""Import DOCX do dokumentu transkrypcji („najlepszy wysiłek”, bez Qt).

Odczytuje runy python-docx i buduje tekst z zakresami indeksu górnego.
Akapity DOCX stają się akapitami dokumentu (rozdzielonymi znakiem nowej
linii). Pogrubiony prefiks ``B:`` na początku akapitu (eksport aplikacji)
przypisuje akapit mówcy, a wiersz „Mówcy” metryczki daje opisy mówców.
Metadane pochodzą z tabeli metryczki (rozpoznawanej po etykietach pól),
akapitu „Autor: … Data: …” z eksportu aplikacji oraz - gdy tych brak -
z ``core_properties``.

Wszystko, czego dokument aplikacji nie wyraża (formatowanie inne niż indeks
górny, tabele poza metryczką, obrazy, przypisy), jest pomijane i liczone
w raporcie.
"""

from __future__ import annotations

import re
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as DocxDocument
from docx.styles.style import CharacterStyle
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from transkryptor.document.metadata import (
    DEFAULT_FIELDS,
    PLACE,
    MetadataField,
    label_variants,
)
from transkryptor.document.model import normalize_ranges
from transkryptor.document.speakers import (
    DEFAULT_SPEAKERS,
    Speaker,
    guess_role,
    parse_legend,
)
from transkryptor.errors import ImportDocxError
from transkryptor.i18n import all_translations, tr

# Pogrubiona etykieta mówcy na początku akapitu: ``B:`` (+ tabulator/spacja).
_SPEAKER_PREFIX_RE = re.compile(r"^\s*(\w{1,8}):\s*")
# Style znakowe runów importowanego dokumentu według identyfikatora stylu.
_StyleCache = dict[str | None, CharacterStyle | None]


def _header_re() -> re.Pattern[str]:
    """Akapit „Autor: … Data: …” w dowolnym języku eksportu."""
    author = "|".join(map(re.escape, all_translations("docx.author")))
    date = "|".join(map(re.escape, all_translations("docx.date")))
    return re.compile(
        rf"^\s*(?:{author}):\s*(?P<author>.*?)\s*(?:{date}):\s*(?P<date>.*?)\s*$"
    )


@dataclass(frozen=True)
class ImportReport:
    """Co pominięto przy imporcie (liczniki wystąpień)."""

    formatted_runs: dict[str, int] = field(default_factory=dict)
    skipped_tables: int = 0
    images: int = 0
    footnotes: int = 0
    unknown_metadata: tuple[str, ...] = ()

    @property
    def is_lossless(self) -> bool:
        return not self.lines()

    def lines(self) -> list[str]:
        """Pozycje raportu dla użytkownika (pusta lista = nic nie pominięto)."""
        lines: list[str] = []
        if self.formatted_runs:
            kinds = ", ".join(
                f"{name} ({count})" for name, count in self.formatted_runs.items()
            )
            lines.append(tr("docx_import.report.formatting", kinds=kinds))
        if self.skipped_tables:
            lines.append(tr("docx_import.report.tables", count=self.skipped_tables))
        if self.images:
            lines.append(tr("docx_import.report.images", count=self.images))
        if self.footnotes:
            lines.append(tr("docx_import.report.footnotes", count=self.footnotes))
        if self.unknown_metadata:
            lines.append(
                tr(
                    "docx_import.report.unknown_metadata",
                    fields=", ".join(self.unknown_metadata),
                )
            )
        return lines


@dataclass(frozen=True)
class ImportedDocument:
    text: str
    superscript_ranges: list[tuple[int, int]]
    author: str
    date: str
    metadata: dict[str, str]
    report: ImportReport
    speakers: tuple[Speaker, ...] = DEFAULT_SPEAKERS
    paragraph_speakers: dict[int, str] = field(default_factory=dict)


def import_docx(
    path: str | Path, fields: Iterable[MetadataField] = DEFAULT_FIELDS
) -> ImportedDocument:
    """Wczytuje DOCX; nieczytelny plik → ``ImportDocxError``."""
    source = Path(path)
    try:
        docx = DocxDocument(str(source))
    except Exception as exc:  # python-docx zgłasza różne wyjątki dla złego pliku
        raise ImportDocxError(
            user_message=tr("docx_import.error.unreadable", name=source.name),
            retry_hint=tr("docx_import.error.unreadable.hint"),
        ) from exc
    return _Reader(docx, fields).read()


@dataclass
class _Blocks:
    """Wynik pierwszego przejścia po dokumencie: nagłówek i akapity treści."""

    paragraphs: list[Paragraph] = field(default_factory=list)
    metadata: dict[str, str] = field(default_factory=dict)
    unknown: list[str] = field(default_factory=list)
    author: str | None = None
    date: str = ""
    skipped_tables: int = 0


class _Reader:
    def __init__(self, docx, fields: Iterable[MetadataField]) -> None:
        self._docx = docx
        labels: dict[str, str] = {}
        for metadata_field in (*DEFAULT_FIELDS, *fields):
            for label in label_variants(metadata_field):
                labels[_label_key(label)] = metadata_field.key
        self._labels = labels
        self._header_re = _header_re()
        self._speaker_labels = {
            _label_key(label) for label in all_translations("docx.speakers")
        }
        self._legend: list[Speaker] = []
        self._formatted: dict[str, int] = {}
        self._styles: _StyleCache = {}

    def read(self) -> ImportedDocument:
        head = self._read_blocks()
        text, ranges, assignments = self._body_text(head.paragraphs)
        core = self._docx.core_properties
        author = head.author if head.author is not None else core.author or ""
        if PLACE not in head.metadata and core.keywords:
            head.metadata[PLACE] = core.keywords
        speakers = self._speakers(assignments)
        body = self._docx.element.body
        return ImportedDocument(
            text=text,
            superscript_ranges=normalize_ranges(ranges),
            author=author,
            date=head.date,
            metadata=head.metadata,
            speakers=tuple(speakers) if speakers else DEFAULT_SPEAKERS,
            paragraph_speakers=assignments,
            report=ImportReport(
                formatted_runs=dict(self._formatted),
                skipped_tables=head.skipped_tables,
                images=len(body.xpath(".//w:drawing | .//w:pict")),
                footnotes=len(
                    body.xpath(".//w:footnoteReference | .//w:endnoteReference")
                ),
                unknown_metadata=tuple(head.unknown),
            ),
        )

    def _read_blocks(self) -> _Blocks:
        """Nagłówek dokumentu (metryczka, autor i data) i akapity treści."""
        blocks = _Blocks()
        # Nagłówek (metryczka, „Autor: … Data: …”) kończy pierwszy akapit
        # z tekstem; flaga zamiast przeglądania wszystkich akapitów w pętli.
        seen_text = False
        for block in self._docx.iter_inner_content():
            in_header = not seen_text
            if isinstance(block, Table):
                if not (
                    in_header
                    and self._read_metadata(block, blocks.metadata, blocks.unknown)
                ):
                    blocks.skipped_tables += 1
                continue
            if in_header and blocks.author is None:
                match = self._header_re.match(block.text)
                if match:
                    blocks.author, blocks.date = match["author"], match["date"]
                    continue
            blocks.paragraphs.append(block)
            seen_text = seen_text or bool(block.text.strip())
        paragraphs = blocks.paragraphs
        while paragraphs and not paragraphs[0].text.strip():
            paragraphs.pop(0)
        while paragraphs and not paragraphs[-1].text.strip():
            paragraphs.pop()
        return blocks

    def _body_text(
        self, paragraphs: list[Paragraph]
    ) -> tuple[str, list[tuple[int, int]], dict[int, str]]:
        """Tekst akapitów bez prefiksów mówców, zakresy indeksu górnego i mówcy."""
        text_parts: list[str] = []
        ranges: list[tuple[int, int]] = []
        assignments: dict[int, str] = {}
        offset = 0
        legend_keys = {speaker.key for speaker in self._legend}
        for index, paragraph in enumerate(paragraphs):
            if index:
                text_parts.append("\n")
                offset += 1
            runs = _runs(paragraph)
            key, skip = _speaker_prefix(runs, legend_keys, self._styles)
            if key is not None:
                assignments[index] = key
            for run in runs:
                run_text = run.text
                if skip:
                    cut = min(skip, len(run_text))
                    skip -= cut
                    run_text = run_text[cut:]
                if not run_text:
                    continue
                if _is_superscript(run, self._styles):
                    ranges.append((offset, offset + len(run_text)))
                self._count_formatting(run)
                text_parts.append(run_text)
                offset += len(run_text)
        return "".join(text_parts), ranges, assignments

    def _speakers(self, assignments: dict[int, str]) -> list[Speaker]:
        """Legenda mówców uzupełniona o klucze użyte tylko w akapitach."""
        speakers = list(self._legend)
        known = {speaker.key for speaker in speakers}
        for key in dict.fromkeys(assignments.values()):
            if key not in known:
                speakers.append(Speaker(key, "", guess_role(key)))
                known.add(key)
        return speakers

    def _read_metadata(
        self, table: Table, metadata: dict[str, str], unknown: list[str]
    ) -> bool:
        """Tabela metryczki: dwie kolumny, co najmniej jedna znana etykieta."""
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not rows or any(len(cells) != 2 for cells in rows):
            return False
        if not any(
            _label_key(label) in self._labels
            or _label_key(label) in self._speaker_labels
            for label, _value in rows
        ):
            return False
        for label, value in rows:
            if _label_key(label) in self._speaker_labels:
                self._legend = parse_legend(value)
                continue
            key = self._labels.get(_label_key(label))
            if key is None:
                if label or value:
                    unknown.append(label or value)
            elif value:
                metadata[key] = value
        return True

    def _count_formatting(self, run: Run) -> None:
        font = run.font
        for name, present in (
            ("docx_import.format.bold", font.bold),
            ("docx_import.format.italic", font.italic),
            ("docx_import.format.underline", font.underline),
            ("docx_import.format.strike", font.strike),
            ("docx_import.format.subscript", font.subscript),
        ):
            if present:
                name = tr(name)
                self._formatted[name] = self._formatted.get(name, 0) + 1


def _speaker_prefix(
    runs: list[Run], known: set[str], styles: _StyleCache
) -> tuple[str | None, int]:
    """Klucz mówcy z pogrubionego prefiksu i liczba znaków do pominięcia.

    Prefiks musi być w całości pogrubiony, a po dwukropku stać tabulator
    (tak zapisuje go eksport) - albo klucz musi być w legendzie mówców.
    Pogrubione „Uwaga:” ze spacją nie jest mówcą.
    """
    bold = ""
    for run in runs:
        if not run.text:
            continue
        if not _is_bold(run, styles):
            break
        bold += run.text
    match = _SPEAKER_PREFIX_RE.match(bold)
    if match is None or bold[match.end() :].strip():
        return None, 0
    text = "".join(run.text for run in runs)
    colon = text.index(":") + 1
    if not text.startswith("\t", colon) and match[1] not in known:
        return None, 0
    skip = colon
    while skip < len(text) and text[skip] in " \t":
        skip += 1
    return match[1], skip


def _run_style(run: Run, styles: _StyleCache) -> CharacterStyle | None:
    """Styl znakowy runu, zapamiętany dla identyfikatora stylu.

    ``Run.style`` przy każdym wywołaniu szuka stylu domyślnego wśród
    wszystkich stylów dokumentu - przy tysiącach runów to większość czasu
    importu (PERF-05).
    """
    style_id = run._r.style
    if style_id not in styles:
        styles[style_id] = run.style
    return styles[style_id]


def _is_bold(run: Run, styles: _StyleCache) -> bool:
    if run.font.bold is not None:
        return bool(run.font.bold)
    style = _run_style(run, styles)
    while style is not None:
        if style.font.bold is not None:
            return bool(style.font.bold)
        style = style.base_style
    return False


def _runs(paragraph: Paragraph) -> list[Run]:
    """Runy akapitu, także te wewnątrz hiperłączy."""
    runs: list[Run] = []
    for item in paragraph.iter_inner_content():
        if isinstance(item, Hyperlink):
            runs.extend(item.runs)
        else:
            runs.append(item)
    return runs


def _is_superscript(run: Run, styles: _StyleCache) -> bool:
    if run.font.superscript is not None:
        return bool(run.font.superscript)
    style = _run_style(run, styles)
    while style is not None:
        if style.font.superscript is not None:
            return bool(style.font.superscript)
        style = style.base_style
    return False


def _label_key(label: str) -> str:
    return label.strip().rstrip(":").strip().casefold()
