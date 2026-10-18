"""Szablon eksportu DOCX instytucji: placeholdery ``{{pole}}`` (propozycja 19).

Szablon to zwykły plik ``.docx`` (czcionki, marginesy, nagłówek, stopka,
logo). Placeholdery ``{{nazwa}}`` w akapitach treści, tabelach, polach
tekstowych, nagłówkach i stopkach są wypełniane metryczką nagrania —
także gdy Word rozbił placeholder na kilka runów (zostaje format pierwszego).
Nazwy są porównywane bez wielkości liter i polskich znaków
(``{{Miejscowość}}`` = ``{{miejscowosc}}``).

Placeholdery specjalne: ``{{transkrypcja}}`` — miejsce akapitów tekstu
(bez niego tekst trafia na koniec dokumentu), ``{{metryczka}}`` — tabela
metryczki, ``{{autor}}``, ``{{data}}`` (transkrypcji) i ``{{mowcy}}``
(legenda mówców). Szablon bez placeholderów pól i bez ``{{metryczka}}``
dostaje dotychczasowy układ: tabelę metryczki i akapit „Autor/Data” przed
tekstem. Nieznany placeholder zostaje w dokumencie i trafia do raportu.
"""

from __future__ import annotations

import re
import unicodedata
from collections.abc import Callable, Iterable, Iterator
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from transkryptor.document.metadata import DEFAULT_FIELDS, MetadataField
from transkryptor.i18n import tr

PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")

TRANSCRIPT = "transkrypcja"
METADATA = "metryczka"
AUTHOR = "autor"
DATE = "data"
SPEAKERS = "mowcy"
SPECIAL_ALIASES: dict[str, tuple[str, ...]] = {
    TRANSCRIPT: ("transkrypcja", "transcription", "tekst", "text"),
    METADATA: ("metryczka", "metadata"),
    AUTHOR: ("autor", "author"),
    DATE: ("data", "date"),
    SPEAKERS: ("mowcy", "speakers"),
}
# Krótkie nazwy pól domyślnych (obok klucza i etykiety w każdym języku).
FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "signature": ("sygnatura",),
    "place": ("miejscowosc",),
    "district": ("gmina_powiat", "gmina", "powiat"),
    "informant": ("informator",),
    "informant_birth_year": ("rok_urodzenia",),
    "recording_date": ("data_nagrania",),
    "researcher": ("badacz",),
    "notes": ("uwagi",),
}

SPECIAL = "special"
FIELD = "field"
# Akapity czytane wprost z XML (także z pól tekstowych) nie mają obiektu
# rodzica python-docx; do odczytu i zmiany tekstu runów nie jest potrzebny.
_NO_PARENT: Any = None


def placeholder_key(name: str) -> str:
    """Nazwa porównywalna: małe litery bez znaków diakrytycznych, ``_`` w odstępach."""
    decomposed = unicodedata.normalize("NFKD", name.replace("ł", "l").replace("Ł", "L"))
    plain = "".join(char for char in decomposed if not unicodedata.combining(char))
    return re.sub(r"[^0-9a-z]+", "_", plain.casefold()).strip("_")


def placeholder_names(
    fields: Iterable[MetadataField] = DEFAULT_FIELDS,
) -> dict[str, tuple[str, str]]:
    """Znormalizowana nazwa → (rodzaj, klucz) dla wszystkich znanych placeholderów."""
    names: dict[str, tuple[str, str]] = {}
    for special, aliases in SPECIAL_ALIASES.items():
        for alias in aliases:
            names[alias] = (SPECIAL, special)
    for field in (*DEFAULT_FIELDS, *fields):
        candidates = [field.key, field.label, field.display_label]
        candidates += FIELD_ALIASES.get(field.key, ())
        for candidate in candidates:
            key = placeholder_key(candidate)
            if key:
                names.setdefault(key, (FIELD, field.key))
    return names


def suggested_placeholders(fields: Iterable[MetadataField]) -> list[str]:
    """Placeholdery do podpowiedzi w ustawieniach (po jednym na pole)."""
    result = [f"{{{{{name}}}}}" for name in SPECIAL_ALIASES]
    for field in fields:
        alias = FIELD_ALIASES.get(field.key, (placeholder_key(field.label),))[0]
        result.append(f"{{{{{alias}}}}}")
    return result


@dataclass(frozen=True)
class TemplateCheck:
    """Wynik sprawdzenia szablonu: błąd otwarcia albo nieznane placeholdery."""

    error: str | None = None
    unknown: tuple[str, ...] = ()
    placeholders: tuple[str, ...] = ()

    @property
    def ok(self) -> bool:
        return self.error is None


def open_template(path: str | Path) -> Any:
    """Szablon jako dokument python-docx; ``ValueError`` z opisem przy błędzie."""
    source = Path(path)
    if not source.is_file():
        raise ValueError(tr("template.error.missing", name=source.name))
    try:
        return DocxDocument(str(source))
    except Exception as exc:  # python-docx zgłasza różne wyjątki dla złego pliku
        raise ValueError(tr("template.error.unreadable", name=source.name)) from exc


def check_template(
    path: str | Path, fields: Iterable[MetadataField] = DEFAULT_FIELDS
) -> TemplateCheck:
    """ACC-56: czy plik otwiera się w python-docx; nieznane placeholdery."""
    try:
        docx = open_template(path)
    except ValueError as error:
        return TemplateCheck(error=str(error))
    known = placeholder_names(fields)
    found: list[str] = []
    for paragraph in iter_paragraphs(docx):
        for match in PLACEHOLDER_RE.finditer(paragraph_text(paragraph)):
            if match[1] not in found:
                found.append(match[1])
    unknown = tuple(name for name in found if placeholder_key(name) not in known)
    return TemplateCheck(unknown=unknown, placeholders=tuple(found))


# --- przechodzenie po dokumencie -------------------------------------------------


def _parts(docx: Any) -> Iterator[Any]:
    """Elementy XML treści, nagłówków i stopek (bez tworzenia nowych)."""
    yield docx.element.body
    for section in docx.sections:
        for part in (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            # Nagłówek „jak w poprzedniej sekcji” nie ma własnej treści,
            # a odczyt utworzyłby pustą definicję.
            if not part.is_linked_to_previous:
                yield part._element


def iter_paragraphs(docx: Any) -> Iterator[Paragraph]:
    """Wszystkie akapity: treść, tabele (zagnieżdżone), pola tekstowe, nagłówki."""
    seen: set[int] = set()
    for element in _parts(docx):
        if id(element) in seen:
            continue
        seen.add(id(element))
        for p in list(element.iter(qn("w:p"))):
            yield Paragraph(p, _NO_PARENT)


def paragraph_text(paragraph: Paragraph) -> str:
    return "".join(run.text for run in paragraph.runs)


def replace_placeholders(
    paragraph: Paragraph, resolve: Callable[[str], str | None]
) -> list[str]:
    """Zastępuje placeholdery akapitu; zwraca nazwy, których ``resolve`` nie zna.

    Placeholder może przechodzić przez kilka runów — wynik trafia do runu,
    w którym się zaczyna (z jego formatem), a jego reszta znika z kolejnych.
    """
    runs = paragraph.runs
    texts = [run.text for run in runs]
    text = "".join(texts)
    if "{{" not in text:
        return []
    starts: list[int] = []
    offset = 0
    for piece in texts:
        starts.append(offset)
        offset += len(piece)
    unknown: list[str] = []
    for match in reversed(list(PLACEHOLDER_RE.finditer(text))):
        value = resolve(match[1])
        if value is None:
            unknown.append(match[1])
            continue
        start, end = match.span()
        first = max(i for i, s in enumerate(starts) if s <= start)
        last = max(i for i, s in enumerate(starts) if s < end)
        head = runs[first].text[: start - starts[first]]
        tail = runs[last].text[end - starts[last] :]
        if first == last:
            runs[first].text = head + value + tail
        else:
            runs[first].text = head + value
            for index in range(first + 1, last):
                runs[index].text = ""
            runs[last].text = tail
    return list(reversed(unknown))


def find_placeholder(docx: Any, name: str) -> Paragraph | None:
    """Pierwszy akapit treści z placeholderem specjalnym ``name``."""
    names = placeholder_names(())
    for p in docx.element.body.iter(qn("w:p")):
        paragraph = Paragraph(p, _NO_PARENT)
        for match in PLACEHOLDER_RE.finditer(paragraph_text(paragraph)):
            if names.get(placeholder_key(match[1])) == (SPECIAL, name):
                return paragraph
    return None


def uses_field_placeholders(docx: Any, fields: Iterable[MetadataField]) -> bool:
    names = placeholder_names(fields)
    for paragraph in iter_paragraphs(docx):
        for match in PLACEHOLDER_RE.finditer(paragraph_text(paragraph)):
            kind = names.get(placeholder_key(match[1]))
            if kind is not None and (kind[0] == FIELD or kind[1] == METADATA):
                return True
    return False


def remove_placeholder(paragraph: Paragraph, name: str) -> None:
    """Usuwa placeholder z akapitu; pusty akapit znika z dokumentu."""
    names = placeholder_names(())

    def resolve(found: str) -> str | None:
        return "" if names.get(placeholder_key(found)) == (SPECIAL, name) else None

    replace_placeholders(paragraph, resolve)
    if not paragraph_text(paragraph).strip() and not paragraph._p.xpath(
        ".//w:drawing | .//w:pict"
    ):
        parent = paragraph._p.getparent()
        if parent is not None:
            parent.remove(paragraph._p)


def empty_copy(paragraph: Paragraph) -> Any:
    """Nowy element akapitu z właściwościami (stylem) akapitu wzorcowego."""
    element = deepcopy(paragraph._p)
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)
    return element
